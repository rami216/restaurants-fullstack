#backend/ai/router/.py
import os, json, re, traceback
from uuid import UUID
from fastapi import APIRouter, HTTPException,Depends,Body
from pydantic import BaseModel, Field
from openai import OpenAI, OpenAIError
from typing import Dict, Any, List, Optional
import re
import anthropic

from sqlalchemy import select
from config import AI_DEFAULT_MODEL
from ai.billing import track_ai_usage
from auth.auth_handler import get_current_active_user
from sqlalchemy.ext.asyncio import AsyncSession
from database import get_db
from models import User
from models import CustomDataSchema
from website_builder.custom_data_router import SchemaField
from .utils import get_ai_client
from website_builder.models import Website
from sqlalchemy.orm import selectinload
from sqlalchemy.orm.attributes import flag_modified
from website_builder.models import (
    Page, Section, Subsection, Element, Navbar, NavbarItem
)
from website_builder.models import SchemaAutomation

router = APIRouter(prefix="/ai", tags=["Extras"])
openai = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))
## using: NON_TABLE_COMPRESSED_TRY1, REFINE_MASTER_PROMPT_2,REFINE_SECTION_SYSTEM_PROMPT,page_generator_test_2,NEW_2_DATA_APP_GENERATOR_PROMPT,VIEW_ONLY_GENERATOR_PROMPT,REFINE_DATA_APP_PROMPT
anthropic_client = anthropic.Anthropic(api_key=os.getenv("CLAUDE_API_KEY"))

import asyncio
from .prompt_modules import (
    build_system_prompt,
    lint_component,
    REPAIR_PROMPT,
    REFINE_OPS_BASE,
    apply_ops,
    PAGE_PLAN_PROMPT,
    DATA_APP_PROMPT_V3,
    sanitize_injected_params,   # ← ADD THIS
    inject_runtime_lib,
    APP_ARCHITECT_PROMPT,
    COMPLEX_ARCHITECT_PROMPT,
    COMPLEX_BUILDER_PROMPT,
    COMPLEX_LIB_REGISTRY
)
from website_builder.models import SchemaAutomation


#region appbuilder

class GenerateAppRequest(BaseModel):
    website_id: UUID
    prompt: str
 
 
@router.post("/generate-app")
async def generate_app(
    body: GenerateAppRequest,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_active_user),
):
    '''
    Orchestrated multi-step build: plan -> tables -> automations -> pages/sections/elements.
    Returns a manifest of everything created. The frontend then refetches the website.
    '''
    try:
        website = await get_website(body.website_id, db)
        client, model, provider = get_ai_client(website)
 
        # ---- STEP 1: ARCHITECT ----
        plan = call_ai_json(
            client, model, provider, APP_ARCHITECT_PROMPT, body.prompt,
            temperature=0.4, max_tokens=4000,
        )
        theme = plan.get("theme", {})
        manifest = {"app_name": plan.get("app_name", "New App"),
                    "tables": [], "pages": [], "automations": 0, "errors": []}
 
        # ---- STEP 2: TABLES (two passes so relations can resolve by name) ----
        name_to_schema_id = {}
 
        # pass 2a: create every table WITHOUT relation fields first (so targets exist)
        raw_tables = plan.get("tables", []) or []
        for t in raw_tables:
            if not isinstance(t, dict) or not t.get("name"):
                continue
            name = t["name"]
            existing = (await db.execute(select(CustomDataSchema).where(
                CustomDataSchema.website_id == body.website_id,
                CustomDataSchema.name == name,
            ))).scalars().first()
            if existing:
                name_to_schema_id[name] = str(existing.schema_id)
                continue
            non_rel_fields = [f for f in (t.get("fields") or []) if f.get("type") != "relation"]
            ns = CustomDataSchema(website_id=body.website_id, name=name, fields=non_rel_fields)
            db.add(ns)
            await db.flush()
            name_to_schema_id[name] = str(ns.schema_id)
        await db.commit()
 
        # pass 2b: now resolve relation fields (related_table name -> schema_id) and patch each schema
        for t in raw_tables:
            if not isinstance(t, dict) or not t.get("name"):
                continue
            name = t["name"]
            sid = name_to_schema_id.get(name)
            if not sid:
                continue
            resolved = []
            for f in (t.get("fields") or []):
                fd = dict(f)
                if fd.get("type") == "relation":
                    target = name_to_schema_id.get(fd.pop("related_table", ""))
                    if not target:
                        continue  # skip a relation whose target wasn't created
                    fd["related_schema_id"] = target
                resolved.append(fd)
            schema = await db.get(CustomDataSchema, UUID(sid))
            if schema:
                schema.fields = resolved
                flag_modified(schema, "fields")   # ⚠️ ensure JSON change is detected
            manifest["tables"].append({"name": name, "schema_id": sid})
        await db.commit()
 
        # ---- STEP 3: AUTOMATIONS (resolve table name -> schema_id) ----
        for a in plan.get("automations", []) or []:
            if not isinstance(a, dict):
                continue
            sid = name_to_schema_id.get(a.get("table", ""))
            if not sid:
                continue
            db.add(SchemaAutomation(
                schema_id=UUID(sid),
                trigger=a.get("trigger", "on_create"),
                action_type=a.get("action_type", "mutate_row"),
                config=a.get("config", {}) or {},
            ))
            manifest["automations"] += 1
        await db.commit()
 
        # snapshot of all schemas for injecting into element props (relation labeling needs this)
        all_schemas = (await db.execute(select(CustomDataSchema).where(
            CustomDataSchema.website_id == body.website_id))).scalars().all()
        all_schemas_summary = [
            {"name": s.name, "schema_id": str(s.schema_id), "fields": s.fields} for s in all_schemas
        ]
 
        # ---- STEP 4: PAGES via the full page engine + data-apps as extra sections ----
        for p_idx, page_plan in enumerate(plan.get("pages", []) or []):
            if not isinstance(page_plan, dict):
                continue
            slug = page_plan.get("slug") or (f"/{page_plan.get('title','page').lower().replace(' ','-')}")

            page = (await db.execute(select(Page).where(
                Page.website_id == body.website_id, Page.slug == slug))).scalars().first()
            if not page:
                page = Page(title=page_plan.get("title", "Page"), slug=slug,
                            website_id=body.website_id, properties={})
                db.add(page)
                await db.flush()
                if slug != "/" and page_plan.get("in_navbar", True):
                    nav = (await db.execute(
                        select(Navbar).options(selectinload(Navbar.items))
                        .where(Navbar.website_id == body.website_id)
                    )).scalars().first()
                    if nav and not any(i.link_url == slug for i in nav.items):
                        db.add(NavbarItem(navbar_id=nav.navbar_id, text=page.title,
                                          link_url=slug, position=len(nav.items) + 1))
                await db.commit()
                await db.refresh(page)

            page_manifest = {"title": page.title, "slug": slug, "sections": 0, "data_apps": 0}
            briefs = page_plan.get("sections", []) or []
            visual_briefs = [b for b in briefs if b.get("element_kind") != "data_app"]
            data_briefs = [b for b in briefs if b.get("element_kind") == "data_app"]

            # 4a. VISUAL sections through the full page engine (rich prompt from the plan)
            try:
                page_prompt = (
                    f"A page titled '{page_plan.get('title','Page')}' for the app "
                    f"'{plan.get('app_name','')}'. It must contain these sections, "
                    f"in this order and spirit:\n"
                    + "\n".join(f"- {b.get('section_type','content')}: {b.get('description','')}"
                                for b in visual_briefs)
                    + "+ \n EXISTING TABLES you may bind sections to (use the exact name as data_binding): "
                    + ", ".join(app_tables.keys())
                    + ". Declare data_tables: [] — never redefine these."
                )
                app_tables = {s["name"]: {"schema_id": s["schema_id"], "fields": s["fields"]}
                              for s in all_schemas_summary}
                sections, _ = await build_page_sections(
                    client, model, provider, page_prompt, db,
                    body.website_id, website.subdomain, theme=theme,
                    preexisting_tables=app_tables,
                )
            except Exception as e:
                manifest["errors"].append(f"{slug} page engine: {e}")
                sections = []

            pos = 1
            for spayload in sections:
                el_payloads = extract_element_payloads(spayload)
                if not el_payloads:
                    manifest["errors"].append(
                        f"{slug}: section had no extractable elements (keys: {list(spayload.keys()) if isinstance(spayload, dict) else type(spayload)})"
                    )
                    continue
                section = Section(page_id=page.page_id, section_type="content",
                                  position=pos, properties={})
                db.add(section)
                await db.flush()
                subsection = Subsection(section_id=section.section_id, position=1,
                                        properties={"flexDirection": spayload.get("layout", "column") if isinstance(spayload, dict) else "column",
                                                    "alignItems": "center"})
                db.add(subsection)
                await db.flush()
                for e_idx, ep in enumerate(el_payloads):
                    props = ep.get("properties", {}) or {}
                    props.setdefault("subdomain", website.subdomain)
                    db.add(Element(
                        subsection_id=subsection.subsection_id, element_type="AI",
                        position=e_idx + 1, properties=props,
                        ai_payload={
                            "aiTemplate": ep.get("aiTemplate", ""),
                            "properties": props,
                            "editableProps": ep.get("editableProps", []),
                            "script": ep.get("script", ""),   # engine already sanitized+lib'd nested scripts
                        },
                    ))
                pos += 1
                page_manifest["sections"] += 1

            # 4b. DATA-APP sections via the specialist generator, appended after the visuals
            for brief in data_briefs:
                try:
                    binding_name = brief.get("data_binding")
                    sid = name_to_schema_id.get(binding_name)
                    if not sid:
                        manifest["errors"].append(f"{slug}: data_binding '{binding_name}' not found")
                        continue
                    da_content = (
                        f'PROMPT: "{brief.get("element_prompt", brief.get("description",""))}"\n\n'
                        f'UNIQUE_CLASS_NAME: `.app-el-{p_idx}-{pos}`\n\n'
                        f'EXISTING_SCHEMAS_ON_WEBSITE: {json.dumps(all_schemas_summary)}\n\n'
                        f'BIND_TO_EXISTING_SCHEMA_ID: {sid}'
                    )
                    payload = call_ai_json(client, model, provider, DATA_APP_PROMPT_V3,
                                           da_content, temperature=0.5)
                    payload = clean_script(payload)
                    props = payload.get("properties", {}) or {}
                    props.update({
                        "schema_id": sid, "originalType": "DATA_TABLE",
                        "schema_fields": next((s["fields"] for s in all_schemas_summary
                                               if s["name"] == binding_name), []),
                        "all_schemas": all_schemas_summary,
                        "website_id": str(body.website_id),
                        "subdomain": website.subdomain,
                    })
                    for auto in payload.get("automations", []) or []:
                        if isinstance(auto, dict):
                            db.add(SchemaAutomation(schema_id=UUID(sid),
                                                    trigger=auto.get("trigger", "on_create"),
                                                    action_type=auto.get("action_type", "mutate_row"),
                                                    config=auto.get("config", {}) or {}))
                    section = Section(page_id=page.page_id, section_type="data_app",
                                      position=pos, properties={})
                    db.add(section)
                    await db.flush()
                    subsection = Subsection(section_id=section.section_id, position=1,
                                            properties={"flexDirection": "column", "alignItems": "center"})
                    db.add(subsection)
                    await db.flush()
                    db.add(Element(
                        subsection_id=subsection.subsection_id, element_type="AI", position=1,
                        properties=props,
                        ai_payload={
                            "aiTemplate": payload.get("aiTemplate", ""),
                            "properties": props,
                            "editableProps": payload.get("editableProps", []),
                            "script": inject_runtime_lib(sanitize_injected_params(payload.get("script", ""))),
                        },
                    ))
                    pos += 1
                    page_manifest["data_apps"] += 1
                except Exception as e:
                    manifest["errors"].append(f"{slug} data-app: {e}")

            await db.commit()
            manifest["pages"].append(page_manifest)
 
        return manifest
 
    except HTTPException:
        raise
    except Exception as e:
        import traceback; traceback.print_exc()
        raise HTTPException(500, f"generate-app failed: {e}")

#endregion
#region helpers



async def get_website(website_id, db: AsyncSession) -> Website:
    result = await db.execute(select(Website).where(Website.website_id == website_id))
    website = result.scalars().first()
    if not website:
        raise HTTPException(404, "Website not found")
    return website
 
 
def _is_newer_openai(model: str) -> bool:
    """Newer OpenAI models (gpt-5.x, o-series, luna...) renamed max_tokens →
    max_completion_tokens and only accept the default temperature."""
    m = str(model).lower()
    return any(t in m for t in ["gpt-5", "gpt-6", "o1", "o3", "o4", "luna"])


def _is_reasoning_model(model: str) -> bool:
    m = str(model).lower()
    return any(t in m for t in ["gpt-5", "gpt-6", "o1", "o3", "o4", "luna"])


def _openai_kwargs(model, system_prompt, user_content, temperature, max_tokens, json_mode):
    kwargs = {
        "model": model,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_content},
        ],
    }
    if json_mode:
        kwargs["response_format"] = {"type": "json_object"}
    if max_tokens:
        kwargs["max_tokens"] = max_tokens
    kwargs["temperature"] = temperature
    return kwargs


def _extract_json(content: str) -> dict:
    content = content.strip()
    content = re.sub(r"^```json?\s*", "", content)
    content = re.sub(r"\s*```$", "", content)
    try:
        return json.loads(content)
    except json.JSONDecodeError:
        pass
    m = re.search(r"\{[\s\S]*\}", content)
    if not m:
        raise HTTPException(500, "Model returned no valid JSON (possibly truncated — raise max_tokens).")
    try:
        return json.loads(m.group(0))
    except json.JSONDecodeError as e:
        raise HTTPException(500, f"Model returned malformed JSON (likely truncated — raise max_tokens): {e}")


def call_ai_json(client, model: str, provider: str, system_prompt: str,
                 user_content: str, temperature: float = 0.2,
                 max_tokens: int = 4096) -> dict:
    if provider == "claude":
        resp = client.messages.create(
            model=model,
            max_tokens=max_tokens or 4096,
            temperature=temperature,
            system=[{"type": "text", "text": system_prompt,
                     "cache_control": {"type": "ephemeral"}}],
            messages=[{"role": "user", "content": user_content}],
        )
        return _extract_json(resp.content[0].text)

    if _is_reasoning_model(model):
        # Reasoning models use the Responses API
        resp = client.responses.create(
            model=model,
            instructions=system_prompt,
            input=user_content,
            reasoning={"effort": "medium"},
            max_output_tokens=max_tokens or 8000,
        )
        return _extract_json(resp.output_text or "")

    resp = client.chat.completions.create(
        **_openai_kwargs(model, system_prompt, user_content, temperature, max_tokens, json_mode=True)
    )
    return _extract_json(resp.choices[0].message.content or "")


def call_ai_text(client, model: str, provider: str, system_prompt: str,
                 user_content: str, temperature: float = 0.2,
                 max_tokens: int = 4096) -> str:
    if provider == "claude":
        resp = client.messages.create(
            model=model,
            max_tokens=max_tokens or 4096,
            temperature=temperature,
            system=system_prompt,
            messages=[{"role": "user", "content": user_content}],
        )
        return resp.content[0].text.strip()

    if _is_reasoning_model(model):
        resp = client.responses.create(
            model=model,
            instructions=system_prompt,
            input=user_content,
            reasoning={"effort": "medium"},
            max_output_tokens=max_tokens or 8000,
        )
        return (resp.output_text or "").strip()

    resp = client.chat.completions.create(
        **_openai_kwargs(model, system_prompt, user_content, temperature, max_tokens, json_mode=False)
    )
    return (resp.choices[0].message.content or "").strip()
 
def clean_script(payload: dict) -> dict:
    if isinstance(payload.get("script"), str):
        m = re.search(r"<script.*?>([\s\S]*?)</script>", payload["script"])
        if m:
            payload["script"] = m.group(1).strip()
    return payload

def generate_with_repair(client, model, provider, system_prompt, user_content, max_repairs=1, max_tokens=4096):
    payload = clean_script(call_ai_json(client, model, provider, system_prompt, user_content, max_tokens=max_tokens))
    for _ in range(max_repairs):
        errors = lint_component(payload)
        if not errors:
            break
        payload = clean_script(call_ai_json(
            client, model, provider, REPAIR_PROMPT,
            "ERRORS TO FIX:\n- " + "\n- ".join(errors)
            + "\n\nCOMPONENT JSON:\n" + json.dumps(payload),
            max_tokens=max_tokens,
        ))
    return payload
 
def inject_schemas(payload: dict, existing_schemas, website_id) -> dict:
    all_schemas = [
        {"name": s.name, "schema_id": str(s.schema_id), "fields": s.fields}
        for s in existing_schemas
    ]
    if "properties" not in payload:
        payload["properties"] = {}
    payload["properties"]["website_id"] = str(website_id)
    payload["properties"]["all_schemas"] = all_schemas
    if "schema_id" in payload.get("properties", {}):
        sid = payload["properties"]["schema_id"]
        match = next((s for s in existing_schemas if str(s.schema_id) == sid), None)
        if match:
            payload["properties"]["schema_fields"] = match.fields
    return payload


#endregion helpers


SECTION_SYSTEM_PROMPT = """
You are an expert layout designer creating the content for a website section. Your task is to generate a valid JSON object representing the 'subsections' and 'elements' based on a user's prompt.

Your output MUST be a valid JSON object containing a single key: "subsections".

**CRITICAL RULES FOR YOUR OUTPUT:**
1.  **Structure:** The value of "subsections" must be an array of subsection objects.
2.  **Subsection Styling:** Each subsection has a "properties" object. To add visual styles like backgrounds, borders, or padding, you MUST put them inside a nested "style" object within "properties".
3.  **Elements:** Each subsection must have an "elements" array containing the content (like "TEXT", "IMAGE", "BUTTON").
4.  **Content:** Fill the elements with relevant placeholder content based on the user's prompt.

**INPUT:** A user's prompt describing the desired section layout.

**OUTPUT:** A valid JSON object containing only the "subsections" array.

**Example Prompt:** "A single subsection with a blue to green gradient background and a welcome title."
**Example Output:**
{
  "subsections": [
    {
      "properties": {
        "display": "flex",
        "flexDirection": "column",
        "alignItems": "center",
        "justifyContent": "center",
        "style": {
          "padding": "4rem",
          "borderRadius": "16px",
          "backgroundImage": "linear-gradient(to right, #3b82f6, #10b981)"
        }
      },
      "elements": [
        {
          "element_type": "TEXT",
          "properties": { "content": "Welcome to Our Website", "style": { "fontSize": "2.5rem", "color": "#FFFFFF" } },
          "aiPayload": null
        }
      ]
    }
  ]
}
""".strip()



#region generate-ai-element
class GenerateRequest(BaseModel):
    prompt: str
    website_id: UUID | str


class GenerateRequestForElement(BaseModel):
    prompt: str
    unique_class_name: str
    website_id: UUID | str

@router.post("/generate-ai-element")
async def generate_ai_element(
    body: GenerateRequestForElement,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_active_user),
):
    try:
        website = await get_website(body.website_id, db)
        client, model, provider = get_ai_client(website)

        schema_result = await db.execute(
            select(CustomDataSchema).where(CustomDataSchema.website_id == body.website_id)
        )
        existing_schemas = schema_result.scalars().all()

        user_content = (
            f'PROMPT: "{body.prompt}"\n\n'
            f'UNIQUE_CLASS_NAME: `.{body.unique_class_name}`\n\n'
            f'EXISTING_SCHEMAS_ON_WEBSITE: {json.dumps([{"name": s.name, "schema_id": str(s.schema_id), "fields": s.fields} for s in existing_schemas])}'
        )

        
        system_prompt = build_system_prompt(body.prompt)
        payload = generate_with_repair(client, model, provider, system_prompt, user_content)
        payload = clean_script(payload)
        payload = inject_schemas(payload, existing_schemas, body.website_id)
        payload["script"] = inject_runtime_lib(sanitize_injected_params(payload.get("script", "")))
        return payload

    except HTTPException:
        raise
    except Exception as e:
        import traceback; traceback.print_exc()
        raise HTTPException(500, f"generate-ai-element failed: {e}")

#endregion generate-ai-element





class RefineStateRequest(BaseModel):
    prompt: str
    currentState: Dict[str, Any]
    website_id: UUID | str





@router.post("/refine-element")
async def refine_element(
    body: RefineStateRequest,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_active_user),
):
    try:
        website = await get_website(body.website_id, db)
        client, model, provider = get_ai_client(website)
 
        # user_content = the request + the full current state.
        # (This was the only thing my earlier note was about — you did it right.)
        user_content = (
            f'USER_PROMPT: "{body.prompt}"\n\n'
            f'CURRENT_COMPONENT_STATE:\n```json\n{json.dumps(body.currentState, indent=2)}\n```'
        )
 
        # Ops-based refine: model returns {"ops":[...]}, server applies them.
        # Untouched keys structurally cannot be lost.
        refine_system = build_system_prompt(body.prompt, base=REFINE_OPS_BASE)
        raw = call_ai_json(client, model, provider, refine_system, user_content)
        # refine may request a new data table (upgrades static elements to functional)
        new_schema_def = raw.get("create_schema")
        if isinstance(new_schema_def, dict) and new_schema_def.get("fields"):
            ns = CustomDataSchema(
                website_id=body.website_id,
                name=new_schema_def.get("name", "New Table"),
                fields=new_schema_def["fields"],
            )
            db.add(ns)
            await db.commit()
            await db.refresh(ns)
            raw.setdefault("ops", []).extend([
                {"op": "set", "path": "properties.schema_id", "value": str(ns.schema_id)},
                {"op": "set", "path": "properties.schema_fields", "value": new_schema_def["fields"]},
            ])
        payload = apply_ops(body.currentState, raw.get("ops", []))
 
        # Safety net: strip Mustache conditionals if the model snuck them in
        ai_template = payload.get("aiTemplate", "")
        if any(tag in ai_template for tag in ["{{#if", "{{#eq", "{{#unless"]):
            for tag in [r'\{\{#if.*?\}\}', r'\{\{#eq.*?\}\}', r'\{\{#unless.*?\}\}',
                        r'\{\{/if\}\}', r'\{\{/eq\}\}', r'\{\{/unless\}\}']:
                ai_template = re.sub(tag, '', ai_template)
            payload["aiTemplate"] = ai_template
 
        payload = clean_script(payload)
        payload["script"] = inject_runtime_lib(sanitize_injected_params(payload.get("script", "")))
        return payload
 
    except HTTPException:
        raise
    except Exception as e:
        import traceback; traceback.print_exc()
        raise HTTPException(500, f"Refine failed: {e}")

  #endregion
  
  
  #region refinesection
  
  # --- START: REFINE SECTION FEATURE ---

# 1. Pydantic model for the request
class RefineSectionRequest(BaseModel):
    prompt: str
    section_json: Dict[str, Any]
    website_id: UUID | str

# 2. System prompt for the AI
REFINE_SECTION_SYSTEM_PROMPT = """
You are an AI assistant that modifies a complete JSON object for a website section.
Your single most important rule is to **start with the user's provided JSON and return the complete, modified version of it.**

**CRITICAL RULES:**
1.  You will be given the `CURRENT SECTION JSON`. Use it as your starting point.
2.  **DO NOT DELETE ANY DATA** unless the user explicitly asks you to. Preserve all existing keys like `section_id`, `subsections`, `elements`, and their properties.
3.  Apply the user's requested change. For visual styles of the main section, modify the `properties.style` object. For subsections, modify their respective `style` objects.
4.  Your final output **MUST BE THE ENTIRE, COMPLETE, MODIFIED JSON OBJECT** for the section. Do not return partial data.
""".strip()

# 3. The API endpoint function
@router.post("/refine-ai-section")
async def refine_ai_section(
    body: RefineSectionRequest,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_active_user),
):
    try:
        website = await get_website(body.website_id, db)
        client, model, provider = get_ai_client(website)
 
        user_content = (
            f'PROMPT: "{body.prompt}"\n\n'
            f'CURRENT SECTION JSON:\n{json.dumps(body.section_json, indent=2)}'
        )
 
        
        return call_ai_json(client, model, provider, REFINE_SECTION_SYSTEM_PROMPT, user_content, temperature=0.5)
 
    except HTTPException:
        raise
    except Exception as e:
        import traceback; traceback.print_exc()
        raise HTTPException(500, f"refine-ai-section failed: {e}")
 


# --- END: REFINE SECTION FEATURE ---


#region pagegenerator
# --- START: NEW PAGE GENERATION FEATURE ---
def extract_element_payloads(section_payload: dict) -> list:
    """Tolerates every shape the section generator may return:
    {elements:[...]} | {subsections:[{elements:[...]}]} | a single element payload."""
    if not isinstance(section_payload, dict):
        return []
    if isinstance(section_payload.get("elements"), list):
        return [e for e in section_payload["elements"] if isinstance(e, dict)]
    if isinstance(section_payload.get("subsections"), list):
        out = []
        for sub in section_payload["subsections"]:
            if isinstance(sub, dict) and isinstance(sub.get("elements"), list):
                out.extend(e for e in sub["elements"] if isinstance(e, dict))
        return out
    if section_payload.get("aiTemplate"):
        return [section_payload]
    return []

async def build_page_sections(client, model, provider, page_prompt, db, website_id, subdomain, theme=None, preexisting_tables=None):
    """The 'good' page generator as a reusable engine: plan → create/reuse declared
    tables → parallel theme-consistent sections → stamp + sanitize + runtime lib.
    Used by /generate-ai-page AND the app orchestrator."""
    plan = call_ai_json(client, model, provider, PAGE_PLAN_PROMPT, page_prompt,
                        temperature=0.4, max_tokens=2000)
    theme = theme or plan.get("theme", {})
    briefs = plan.get("sections", []) or []
    if not briefs:
        raise HTTPException(500, "Page planner returned no sections.")

    # create (or reuse by name) tables this page declares
    table_map = dict(preexisting_tables or {})
    for t in plan.get("data_tables", []) or []:
        if not isinstance(t, dict) or not t.get("fields"):
            continue
        name = t.get("name", "Page Data")
        existing = (await db.execute(select(CustomDataSchema).where(
            CustomDataSchema.website_id == website_id,
            CustomDataSchema.name == name,
        ))).scalars().first()
        if existing:
            table_map[name] = {"schema_id": str(existing.schema_id), "fields": existing.fields}
        else:
            ns = CustomDataSchema(website_id=website_id, name=name, fields=t["fields"])
            db.add(ns)
            await db.flush()
            table_map[name] = {"schema_id": str(ns.schema_id), "fields": t["fields"]}
    if table_map:
        await db.commit()

    def _gen_section(brief):
        content = (
            f'PROMPT: "{brief.get("description", "")}"\n'
            f'LAYOUT: {brief.get("layout", "column")}\n'
            f'THEME (use these exact colors/fonts for consistency): {json.dumps(theme)}'
        )
        binding = table_map.get(brief.get("data_binding") or "")
        if binding:
            content += f"\nDATA_TABLE (bind the form to this): {json.dumps(binding)}"
        return call_ai_json(client, model, provider, SECTION_GENERATOR_PROMPT, content,
                            temperature=0.5, max_tokens=4096)

    raw_sections = await asyncio.gather(
        *[asyncio.to_thread(_gen_section, b) for b in briefs]
    )

    cleaned = []
    for s in raw_sections:
        s = clean_script(s)
        def _walk(node):
            if isinstance(node, dict):
                props = node.get("properties")
                if isinstance(props, dict):
                    if "website_id" in props:
                        props["website_id"] = str(website_id)
                    props.setdefault("subdomain", subdomain)
                if isinstance(node.get("script"), str) and node["script"].strip():
                    node["script"] = inject_runtime_lib(sanitize_injected_params(node["script"]))
                for v in node.values():
                    _walk(v)
            elif isinstance(node, list):
                for v in node:
                    _walk(v)
        _walk(s)
        cleaned.append(s)
    return cleaned, theme
@router.post("/generate-ai-page")
async def generate_ai_page(
    body: GenerateRequest,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_active_user),
):
    try:
        website = await get_website(body.website_id, db)
        client, model, provider = get_ai_client(website)
        sections, theme = await build_page_sections(
            client, model, provider, body.prompt, db,
            body.website_id, website.subdomain,
        )
        return {"sections": sections, "theme": theme}
    except HTTPException:
        raise
    except Exception as e:
        import traceback; traceback.print_exc()
        raise HTTPException(500, f"generate-ai-page failed: {e}")# --- END: NEW PAGE GENERATION FEATURE ---


#endregion pagegenerator  

#region generatesection

class GenerateSectionRequest(BaseModel):
    prompt: str
    website_id: UUID | str

@router.post("/generate-ai-section")
async def generate_ai_section(
    body: GenerateRequest,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_active_user),
):
    try:
        website = await get_website(body.website_id, db)
        client, model, provider = get_ai_client(website)
 
        
        return call_ai_json(client, model, provider, SECTION_GENERATOR_PROMPT, body.prompt, temperature=0.5, max_tokens=2048)
 
    except HTTPException:
        raise
    except Exception as e:
        import traceback; traceback.print_exc()
        raise HTTPException(500, f"generate-ai-section failed: {e}")
 

SECTION_GENERATOR_PROMPT = """

You are a Lead UI/UX Designer and Frontend Architect. Your task is to generate the JSON for a **single, high-fidelity website section** based on a user's prompt.

**OUTPUT FORMAT:**
You must return a valid JSON object representing **ONE Section**.
The object MUST have these three keys: `"section_type"`, `"properties"`, and `"subsections"`.
Do NOT wrap this in a "sections" array. Return the single section object directly.

---

### **ARCHITECTURAL RULES (The "Strong" Layout System)**

**1. SECTION STRUCTURE (The Container):**
   - Each object represents a full-width stripe.
   - **`properties`**: MUST contain `display: "flex"`.
   - **`style`**: Background colors/images and padding (e.g., `padding: "4rem 1rem"`).
   - **Layout Logic:**
     - Vertical Stacking: Set `flexDirection: "column"`, `alignItems: "center"`.
     - Side-by-Side: Set `flexDirection: "row"`, `flexWrap: "wrap"`, `justifyContent: "center"`, `alignItems: "center"`, `gap: "4rem"`.

**2. SUBSECTION STRUCTURE (The Columns/Wrappers):**
   - **CRITICAL:** The `subsections` array MUST be present and MUST NOT be empty.
   - **`properties`**: `display: "flex"`, `flexDirection: "column"`, `gap: "1.5rem"`.
   - **Sizing:**
     - Row layout: `width: "45%"`, `minWidth: "300px"` (for responsiveness).
     - Column layout: `width: "100%"`, `maxWidth: "1280px"`, `textAlign: "center"`.

**3. ELEMENT STRUCTURE (Atomic AI Components):**
   - **STRICT TYPE RULE:** You must **ONLY** generate elements with `element_type: "AI"`.
   - **FORBIDDEN TYPES:** Do NOT generate `element_type: "BUTTON"`, `element_type: "TEXT"`, `element_type: "IMAGE"`, or `element_type: "MENU_ITEM"`. Use the AI wrapper for everything.
   - **CRITICAL:** The `elements` array inside a subsection MUST be present and MUST NOT be empty.
   - **Granularity:** Break content down. One element per Headline, one for Subtitle, one for Button.
   - **JSON Structure (MANDATORY):**
     ```json
     {
       "element_type": "AI",
       "aiPayload": {
         "aiTemplate": "<div class='unique-class'> ... content ... </div>",
         "properties": { "text": "Example", ... },
         "editableProps": [ ... ],
         "script": ""
       }
     }
     ```
   - **DATA SAFETY RULE:** The `properties` key inside `aiPayload` is **MANDATORY**. It cannot be null. It must be an object, even if empty (e.g., `"properties": {}`).

---

### **COMPONENT GENERATION RULES**

**A. HTML & CSS:**
   - Wrap everything in a single `<div>` with a unique class name.
   - Use `<style>` inside `aiTemplate` for all CSS.
   - **SCOPING:** Prefix selectors (e.g., `.ai-card-123 h3`).
   - **STYLING:** Use modern CSS (Flexbox, Grid, Shadows, Rounded Corners).

**B. EDITABILITY:**
   - Replace text/colors with Mustache tokens (e.g., `{{title}}`).
   - Create corresponding keys in `properties` and `editableProps`.
   - **IMAGE RULE:** If using an image, set type to `"image"` in `editableProps`.

**C. INTERACTIVITY:**
   - No forms/databases unless explicitly requested. Use visual elements only.
   
## GRANULARITY (non-negotiable)
- Every distinct link gets its OWN URL token: social icons → facebookUrl, twitterUrl, instagramUrl in properties AND editableProps (type "text"), used as href="{{facebookUrl}}" etc. NEVER href="#", never one shared link for several targets.
- Every button gets {{...Text}} + {{...Url}} tokens. Nav-like link lists: one token pair per link.
- Split into separate elements anything the owner will edit/move independently (each feature card's image, a CTA button vs its heading).

## FUNCTIONAL WIDGETS (newsletter / subscribe / contact capture must WORK, not just render)
Include a script that: reads the input, disables the button with a loading label, then:
await api.post('/builder/form-submissions', {
  website_id: properties.website_id,
  form_element_id: properties.form_id,
  submission_data: { email: emailInput.value }
});
On success replace the form with a thank-you line ({{successMessage}} token); on error alert(err.response?.data?.detail || 'Something went wrong').
Requirements: properties must include "website_id": "WEBSITE_ID_PLACEHOLDER" and "form_id": a literal random UUID string you generate (e.g. "a3f1c2d4-5b6e-4f7a-8c9d-0e1f2a3b4c5d"); button uses btn.onclick with e.preventDefault-safe form; container.querySelector only. The owner sees submissions in their dashboard.
## DATA_TABLE BINDING (when the user content includes a DATA_TABLE json)
Bind the form to it instead of form-submissions: properties.schema_id = its schema_id; inputs named with its exact field ids; submit → await api.post(`/custom-data/rows/${schemaId}`, { data, sitemember_id: null }) (schemaId is injected at runtime from properties.schema_id); on error alert(err.response?.data?.detail || 'Something went wrong') — the server enforces required/unique and returns readable messages (e.g. duplicate newsletter email → 409). Success: in the SCRIPT do msgEl.textContent = properties.successMessage || 'Thanks for subscribing!' — never write a literal {{token}} in JS (it prints as-is); {{successMessage}} belongs only in aiTemplate if pre-placed as a hidden element.
## ENTRANCE ANIMATIONS
Give content blocks class "reveal" with the scroll-gated pattern (.reveal hidden → .visible via onVisible, prefers-reduced-motion respected) — onVisible is pre-injected. Below-the-fold sections must NOT animate on load. MANDATORY: every non-hero section's main inner wrapper carries class "reveal" (opacity:0/translateY CSS + onVisible → .visible + prefers-reduced-motion guard); only the hero may animate immediately.

---

### **EXAMPLE OUTPUT (A Hero Section):**

```json
{
  "section_type": "hero",
  "properties": {
    "display": "flex",
    "flexDirection": "row",
    "flexWrap": "wrap",
    "justifyContent": "center",
    "alignItems": "center",
    "gap": "4rem",
    "style": { "backgroundColor": "#111827", "padding": "6rem 2rem" }
  },
  "subsections": [
    {
      "properties": { "style": { "width": "45%", "minWidth": "320px", "display": "flex", "flexDirection": "column", "gap": "20px", "alignItems": "flex-start" } },
      "elements": [
        {
          "element_type": "AI",
          "aiPayload": {
            "aiTemplate": "<div class='ai-head-01'><h1 style='font-size:3.5rem; color:{{color}}; margin:0;'>{{text}}</h1></div>",
            "properties": { "text": "Delicious Food", "color": "#ffffff" },
            "editableProps": [ {"key":"text","label":"Text","type":"text"}, {"key":"color","label":"Color","type":"color"} ],
            "script": ""
          }
        },
        {
          "element_type": "AI",
          "aiPayload": {
            "aiTemplate": "<button style='background:{{bg}}; color:{{color}}; padding:12px 32px; border-radius:8px; border:none; cursor:pointer;'>{{label}}</button>",
            "properties": { "label": "Order Now", "bg": "#f59e0b", "color": "#000000" },
            "editableProps": [ {"key":"label","label":"Label","type":"text"}, {"key":"bg","label":"Background","type":"color"} ],
            "script": ""
          }
        }
      ]
    },
    {
      "properties": { "style": { "width": "45%", "minWidth": "320px", "display": "flex", "justifyContent": "center" } },
      "elements": [
        {
          "element_type": "AI",
          "aiPayload": {
            "aiTemplate": "<div class='ai-img-01'><img src='{{src}}' style='width:100%; border-radius:16px;' /></div>",
            "properties": { "src": "[https://placehold.co/600x400](https://placehold.co/600x400)" },
            "editableProps": [ {"key":"src","label":"Image URL","type":"image"} ],
            "script": ""
          }
        }
      ]
    }
  ]
}
INPUT: A user's prompt (e.g., "A pricing section"). OUTPUT: The valid JSON object.
""".strip()
#endregion generatesection

#region data_app_element



# --- THE CORRECTED FUNCTION AND MODELS ---
class GenerateDataAppRequest(BaseModel):
    prompt: str
    website_id: UUID
    unique_class_name: str

class AIResponseSchema(BaseModel):
    name: str
    schema_fields: List[SchemaField] = Field(..., alias="schema")
    ai_template: str = Field(..., alias="aiTemplate")
    properties: Dict[str, Any]
    editable_props: List[Dict[str, Any]] = Field(..., alias="editableProps")
    script: str



#region new generateview

VIEW_ONLY_GENERATOR_PROMPT = """
You are an expert front-end developer creating a READ-ONLY component to display data from an existing data source.

Your output MUST be a valid JSON object with FIVE keys: "name_to_find", "aiTemplate", "properties", "editableProps", and "script".

---
### **CRITICAL RULES FOR YOUR OUTPUT**

1.  **`name_to_find`**: The EXACT name of the data schema to find and display, extracted from the user's prompt (e.g., "User Management System").

2.  **`aiTemplate`**: The main HTML structure. It MUST include:
    -   A `<style>` tag for all CSS.
    -   A static main title `<h3>` or `<h2>`.
    -   An **EMPTY** container for displaying the data (e.g., `<div class="data-display"></div>`).
    -   An EMPTY container for pagination controls (e.g., <div class="pagination-controls"></div>).
    -   A `<template id="displayTemplate">`.
    -   **DO NOT** include an "Add New" button or a form container.

3.  displayTemplate: A Mustache/HTML template for ONE data item.
    - The API sends a row object with this structure: {"row_id": "...", "data": {"field_id": "value"}}.
    - For regular fields, you MUST use {{data.field_id}}.
    - CRITICAL: For relational fields (e.g., a field with id 'project'), the data object will contain a nested object. You MUST access this nested data correctly. Look at the SCHEMA_OF_DATA_TO_DISPLAY context to find the exact field id from the related schema to display (e.g., {{data.project.data.project_title}}).
    - DO NOT include edit or delete buttons.

4.  **Styling & Editable Properties (`properties`, `editableProps`)**:
    -   Make the component's styling fully editable.
    -   All style values and user-facing text (like titles) MUST use mustache tokens.
    -   For EVERY token, add a corresponding entry in `properties` and `editableProps`.
    -   **CRITICAL SCOPING RULE:** Every CSS rule **MUST** be prefixed with the given `unique_class_name`.

5.  **`script`**: A complete, raw JavaScript string that makes the element interactive.
    -   It is executed in a function that receives `(container, api, schemaId, properties, Mustache)`.
    -   State Management: It MUST manage state for currentPage (0-indexed), rowsPerPage (e.g., 20), and totalRows.
    -   It must ONLY fetch and render data.
    -   API Calls to Use:
        -   Fetch Paginated Rows: api.get(/custom-data/rows/schemaId?skip={currentPage * rowsPerPage}&limit=${rowsPerPage}). The response is { "rows": [], "total": 0 }.
    -   Pagination Logic:
            - It MUST render "Previous" and "Next" buttons inside the .pagination-controls container.
            - Buttons MUST be disabled when on the first or last page.
            - Clicking the buttons MUST update the currentPage state and re-fetch the data.
    -   It MUST use function expressions (e.g., `const myFunc = () => {}`).

---
**INPUT:** A user's prompt and a `unique_class_name`.
**OUTPUT:** A single, valid JSON object.

**Example Prompt:** "Show a list of our contacts."
**Example `unique_class_name`:** `.ai-contact-view-12345`
**Example Output:**
{
  "name_to_find": "Contact List",
  "aiTemplate": "<style>.ai-contact-view-12345 h3 { color: {{titleColor}}; } .ai-contact-view-12345 .pagination-controls button { margin: 0 5px; }</style><h3>{{title}}</h3><div class=\\"data-display\\"></div><div class=\\"pagination-controls\\"></div><template id=\\"displayTemplate\\"><div><span><strong>{{data.name}}</strong> ({{data.email}})</span></div></template>",
  "properties": { "title": "Our Contacts", "titleColor": "#333333" },
  "editableProps": [ { "key": "title", "label": "Title", "type": "text" }, { "key": "titleColor", "label": "Title Color", "type": "color" } ],
  "script": "const dataDisplay = container.querySelector('.data-display'); const paginationControls = container.querySelector('.pagination-controls'); const displayTemplate = container.querySelector('#displayTemplate').innerHTML; let currentPage = 0; const rowsPerPage = 20; let totalRows = 0; const fetchAndRenderRows = async () => { try { const response = await api.get(`/custom-data/rows/${schemaId}?skip=${currentPage * rowsPerPage}&limit=${rowsPerPage}`); const { rows, total } = response.data; totalRows = total; dataDisplay.innerHTML = ''; rows.forEach(row => { const div = document.createElement('div'); div.innerHTML = Mustache.render(displayTemplate, row); dataDisplay.appendChild(div); }); renderPagination(); } catch (err) { console.error('Failed to fetch data:', err); } }; const renderPagination = () => { paginationControls.innerHTML = ''; const totalPages = Math.ceil(totalRows / rowsPerPage); if (totalPages <= 1) return; const prevButton = document.createElement('button'); prevButton.textContent = 'Previous'; prevButton.disabled = currentPage === 0; prevButton.addEventListener('click', () => { if (currentPage > 0) { currentPage--; fetchAndRenderRows(); } }); const nextButton = document.createElement('button'); nextButton.textContent = 'Next'; nextButton.disabled = currentPage >= totalPages - 1; nextButton.addEventListener('click', () => { if (currentPage < totalPages - 1) { currentPage++; fetchAndRenderRows(); } }); paginationControls.appendChild(prevButton); paginationControls.appendChild(nextButton); }; fetchAndRenderRows();"
}

""".strip()

# --- NEW PYDANTIC MODELS AND ENDPOINT FOR VIEW-ONLY ---
class GenerateViewOnlyRequest(BaseModel):
    prompt: str
    website_id: UUID
    unique_class_name: str

class AIViewOnlyResponseSchema(BaseModel):
    name_to_find: str
    ai_template: str = Field(..., alias="aiTemplate")
    properties: Dict[str, Any]
    editable_props: List[Dict[str, Any]] = Field(..., alias="editableProps")
    script: str

@router.post("/generate-view-only-element")
async def generate_view_only_element(
    body: GenerateViewOnlyRequest,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_active_user),
):
    try:
        website = await get_website(body.website_id, db)
        client, model, provider = get_ai_client(website)
 
        name_payload = call_ai_json(
            client, model, provider,
            "You are a helpful assistant that extracts information.",
            f'From this prompt, extract the data source name. Respond with JSON key "name_to_find".\n\nPROMPT: "{body.prompt}"',
            temperature=0.0,
        )
        name_to_find = name_payload.get("name_to_find")
        if not name_to_find:
            raise HTTPException(400, "Could not determine data source name from prompt.")
 
        result = await db.execute(
            select(CustomDataSchema)
            .where(CustomDataSchema.website_id == body.website_id)
            .where(CustomDataSchema.name == name_to_find)
        )
        existing_schema = result.scalars().first()
        if not existing_schema:
            raise HTTPException(404, f"Data source '{name_to_find}' not found.")
 
        user_content = (
            f'PROMPT: "{body.prompt}"\n\n'
            f'UNIQUE_CLASS_NAME: `.{body.unique_class_name}`\n\n'
            f'SCHEMA_OF_DATA_TO_DISPLAY: {json.dumps(existing_schema.fields)}'
        )
 
        
        payload = call_ai_json(client, model, provider, VIEW_ONLY_GENERATOR_PROMPT, user_content, temperature=0.5)
        payload = clean_script(payload)
 
        final_props = payload.get("properties", {})
        final_props["schema_id"] = str(existing_schema.schema_id)
        final_props["originalType"] = "DATA_VIEW"
        final_props["website_id"] = str(body.website_id)
 
        return {
            "aiTemplate": f'<div class="{body.unique_class_name}">{payload["aiTemplate"]}</div>',
            "properties": final_props,
            "editableProps": payload.get("editableProps", []),
            "script": payload.get("script", ""),
        }
 
    except HTTPException:
        raise
    except Exception as e:
        import traceback; traceback.print_exc()
        raise HTTPException(500, f"AI View-Only generation failed: {e}")
 






#region refine_data_app
REFINE_DATA_APP_PROMPT = """
You are an expert full-stack developer, and your task is to refine an interactive data-driven component based on user feedback.
You will be given the user's prompt, the current JSON state of the component, and a list of all existing schemas on the website.

Your output MUST be a single, complete, valid JSON object representing the fully updated state.

**CRITICAL RULES:**
1.  **Use Context for Relationships:** When the user's prompt involves a related table (e.g., "show the project's due date"), you **MUST** use the `EXISTING_SCHEMAS_ON_WEBSITE` context to find the exact field IDs of the related schema and modify the `displayTemplate` or `script` accordingly.
2.  **Preserve Core Structure**: DO NOT change the `schema_id` or the `schema_fields` in the `properties` object. The database schema is fixed.
3.  **Apply User's Prompt**:
    -   For visual changes (colors, fonts, layout), modify the CSS in the `aiTemplate` and update the corresponding values in the `properties` object.
    -   For functional changes (e.g., "change the edit button to an icon"), modify the `aiTemplate` and the `script` accordingly.
    -   For text changes (e.g., "change the title to 'Manage Employees'"), update the value in the `properties` object.
4.  **Return the Complete Object**: Your final output must be the entire, valid JSON object for the component, including `aiTemplate`, `properties`, `editableProps`, and `script`.
""".strip()

@router.post("/refine-data-app-element")
async def refine_data_app_element(
    body: RefineStateRequest,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_active_user),
):
    try:
        website = await get_website(body.website_id, db)
        client, model, provider = get_ai_client(website)
 
        schema_result = await db.execute(
            select(CustomDataSchema).where(CustomDataSchema.website_id == body.website_id)
        )
        existing_schemas = schema_result.scalars().all()
 
        user_content = (
            f'USER_PROMPT: "{body.prompt}"\n\n'
            f'CURRENT_COMPONENT_STATE:\n```json\n{json.dumps(body.currentState, indent=2)}\n```\n\n'
            f'EXISTING_SCHEMAS_ON_WEBSITE: {json.dumps([{"name": s.name, "schema_id": str(s.schema_id), "fields": s.fields} for s in existing_schemas])}'
        )
 
        
        return call_ai_json(client, model, provider, REFINE_DATA_APP_PROMPT, user_content)
 
    except HTTPException:
        raise
    except Exception as e:
        import traceback; traceback.print_exc()
        raise HTTPException(500, f"Data App refinement failed: {e}")
 


#region nontabletestingai

#region nontable-compressed
#endregion nontabletestingai

#region ARCHITECT_SYSTEM_PROMPT

class ChatMessage(BaseModel):
    role: str
    content: str
class ArchitectBlueprintRequest(BaseModel):
    idea: str
    history: list[ChatMessage] = [] # <--- Added this
    website_id: UUID

# Paste the massive prompt here so it's safely locked in the backend
ARCHITECT_SYSTEM_PROMPT = """
You are the Master Architect for "Zygoflow", an advanced AI-powered SaaS builder. 
Your job is to listen to a user's app idea and break it down into a highly structured, copy-paste blueprint. 

Zygoflow has two distinct builders:
1. DATA TABLE BUILDER (Creates database schemas and basic CRUD forms).
2. UI BUILDER (Creates complex dashboards, AI generators, workflows, and chatbots).

When the user describes an app, you must analyze the requirements and output a blueprint in this EXACT Markdown format.

---

### 🗄️ Phase 1: Database Tables (Data Table Builder)
List every table needed. For EACH table, provide:
- **Table Name:** [Name]
- **Fields:** [List fields and types. Types must be: text, number, boolean, image, gallery, file, or relation].
- **Table Builder Prompt:** [Write a highly technical, 1-3 sentence prompt for the user to copy/paste].

**CRITICAL TABLE RULES TO INJECT IN THE PROMPT:**
- **Relations:** If a table connects to another, explicitly say: "Create a relational field to the [Target] schema."
- **Visibility:** If it's for admins, say: "Admin mode (fetch and display list on load)." If it's a public form (like booking/contact), say: "Public/Write-Only mode (do not load data, hide form on submit)."
- **Cross-Table Mutations:** If a booking/action changes another table's status, say: "Add a cross-table mutation to set the related [Target] row's [Field] to [Value] on submit."

---

### 🚀 Phase 2: UI Dashboards & Tools (UI Builder)
List the front-end views, workflows, or bots needed. For EACH view, provide:
- **Element Name:** [Name]
- **UI Builder Prompt:** [Write a highly detailed, developer-level prompt for the user to copy/paste].

**CRITICAL UI RULES TO INJECT IN THE PROMPT (Pick the ones that apply):**

**1. AI & Generations:**
- If saving AI data to a database: "Use AI FLAT ARRAY MODE to generate and bulk save data. Include THE SILENCE RULE."
- If showing a complex UI (charts/cards) from AI without saving: "Use AI STRUCTURED UI MODE to extract arrays/scores. Include THE SILENCE RULE."

**2. Chatbots:**
- If the bot just answers questions: "Build a READ-ONLY BOT. Inject database context into the system prompt."
- If the bot books things, saves data, or sends emails: "Build an ACTION-BASED BOT capable of multi-step workflows. Ensure it collects the user email and confirms before executing."

**3. Advanced Integrations:**
- **PDFs:** If analyzing documents: "Use the PDF Upload & AI Extraction chain: Upload -> Parse -> AI -> Save."
- **Emails:** If notifying users: "Integrate email sending via /builder/send-email. Add emailSubject and emailBody to properties."
- **Webhooks:** If sending data to Zapier/Make: "Trigger a Webhook POST request to properties.webhookUrl on success."
- **External APIs:** If fetching 3rd party data: "Use the backend proxy (/builder/fetch-external) and extract from res.data.data."

**4. Data & Logic:**
- **Relations/Dropdowns:** If filtering data (e.g. Times for a Day): "Use cascading dropdowns matching parent row_id to filter children."
- **Privacy:** If the user only sees their own data: "Strictly filter data for the logged-in user using sitemember_id."
- **Math/Stats:** If calculating totals or KPIs: "Use the backend /stats API for math. DO NOT use frontend loops to calculate totals."
- **Payments:** If buying something: "Add an Add-to-Cart system" OR "Add a Checkout Button isolated click listener."

---
**TONE:** Professional, highly structured, acting as a Senior CTO. Give them the exact copy-paste prompts they need to succeed in Zygoflow. Do not explain the code, just give the architecture and the prompts.
"""

@router.post("/generate-architect-blueprint")
async def generate_architect_blueprint(
    body: ArchitectBlueprintRequest,
    db: AsyncSession = Depends(get_db),
    user = Depends(get_current_active_user)
):
    try:
        # 1. Build the message array for OpenAI
        openai_messages = [{"role": "system", "content": ARCHITECT_SYSTEM_PROMPT}]
        
        for msg in body.history[-10:]:
            role = "assistant" if msg.role == "ai" else "user"
            openai_messages.append({"role": role, "content": msg.content})
            
        openai_messages.append({"role": "user", "content": body.idea})

        # 2. Call OpenAI
        resp = openai.chat.completions.create(
            model=AI_DEFAULT_MODEL,
            messages=openai_messages,
            temperature=0.7,
            max_tokens=2500,
        )
        
        # 3. Extract the response correctly
        generated_blueprint = resp.choices[0].message.content

        # 4. Track usage
        usage = getattr(resp, "usage", None)
        prompt_tokens = int(getattr(usage, "prompt_tokens", 0) or 0)
        completion_tokens = int(getattr(usage, "completion_tokens", 0) or 0)
        model_used = getattr(resp, "model", AI_DEFAULT_MODEL)
        
        await track_ai_usage(
            db=db,
            website_id=body.website_id,
            user_id=user.id,
            model=model_used,
            feature="architect_blueprint",
            prompt_tokens=prompt_tokens,
            completion_tokens=completion_tokens,
            meta={"prompt_len": len(body.idea)}
        )

        # 5. Return the result
        return {"result": generated_blueprint}

    except Exception as e:
        import traceback; traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Architect generation failed: {e}")


@router.post("/generate-data-app-element")
async def generate_data_app_element(
    body: GenerateDataAppRequest,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_active_user),
):
    try:
        website = await get_website(body.website_id, db)
        client, model, provider = get_ai_client(website)

        schema_result = await db.execute(
            select(CustomDataSchema).where(CustomDataSchema.website_id == body.website_id)
        )
        existing_schemas = schema_result.scalars().all()

        user_content = (
            f'PROMPT: "{body.prompt}"\n\n'
            f'UNIQUE_CLASS_NAME: `.{body.unique_class_name}`\n\n'
            f'EXISTING_SCHEMAS_ON_WEBSITE: {json.dumps([{"name": s.name, "schema_id": str(s.schema_id), "fields": s.fields} for s in existing_schemas])}'
        )

        # V3 returns ONE flat object:
        # {name, schema, aiTemplate, properties, editableProps, script,
        #  automations, schemas_to_create (optional helpers)}
        payload = generate_with_repair(client, model, provider, DATA_APP_PROMPT_V3, user_content)

        main_fields = payload.get("schema", []) or []
        if not main_fields or "script" not in payload:
            raise HTTPException(500, "AI response missing required structure.")

        # ---- 1. Create helper schemas first (placeholder resolution) ----
        created_schemas_map = {}
        for extra in payload.get("schemas_to_create", []) or []:
            if not isinstance(extra, dict):
                continue
            extra_fields = extra.get("fields") or extra.get("schema") or extra.get("schema_fields") or []
            if not extra_fields:
                continue
            helper_schema = CustomDataSchema(
                website_id=body.website_id,
                name=extra.get("name", "Related Data"),
                fields=extra_fields,
            )
            db.add(helper_schema)
            await db.flush()
            created_schemas_map[f'PLACEHOLDER_FOR_{extra.get("name", "")}'] = str(helper_schema.schema_id)

        # ---- 2. Resolve placeholders + sanitize UUIDs in the main schema ----
        sanitized_fields = []
        for field in main_fields:
            fd = dict(field)
            rel = fd.get("related_schema_id")
            if isinstance(rel, str) and rel in created_schemas_map:
                fd["related_schema_id"] = created_schemas_map[rel]
            elif isinstance(rel, UUID):
                fd["related_schema_id"] = str(rel)
            sanitized_fields.append(fd)

        # ---- 3. Create the main schema ----
        new_schema = CustomDataSchema(
            website_id=body.website_id,
            name=payload.get("name", "Data App"),
            fields=sanitized_fields,
        )
        db.add(new_schema)
        await db.flush()
        final_schema_id = str(new_schema.schema_id)

        # ---- 4. Save declarative automations (server-side booking locks,
        #         stock counters — replaces crossTableMutations JS) ----
        for auto in payload.get("automations", []) or []:
            if not isinstance(auto, dict):
                continue
            db.add(SchemaAutomation(
                schema_id=new_schema.schema_id,
                trigger=auto.get("trigger", "on_create"),
                action_type=auto.get("action_type", "mutate_row"),
                config=auto.get("config", {}) or {},
            ))

        await db.commit()

        # ---- 5. Assemble final props (your existing conventions) ----
        all_schemas_result = await db.execute(
            select(CustomDataSchema).where(CustomDataSchema.website_id == body.website_id)
        )
        all_schemas = all_schemas_result.scalars().all()

        final_props = payload.get("properties", {}) or {}
        final_props["schema_id"] = final_schema_id
        final_props["originalType"] = "DATA_TABLE"
        final_props["schema_fields"] = sanitized_fields
        final_props["all_schemas"] = [{"name": s.name, "schema_id": str(s.schema_id), "fields": s.fields} for s in all_schemas]
        final_props["website_id"] = str(body.website_id)
        final_props["subdomain"] = website.subdomain  # used by user-scoped scripts (localStorage key)

        return {
            "aiTemplate": f'<div class="{body.unique_class_name}">{payload["aiTemplate"]}</div>',
            "properties": final_props,
            "editableProps": payload.get("editableProps", []),
            "script": inject_runtime_lib(sanitize_injected_params(payload["script"])),
        }

    except HTTPException:
        raise
    except Exception as e:
        import traceback; traceback.print_exc()
        raise HTTPException(500, f"AI Data App generation failed: {e}") 
#region agentai-desktop



# Ensure your OpenAI key is loaded

class AnalyzeTextRequest(BaseModel):
    website_id: str  # Mandatory for tracking!
    text: str
    instruction: str    


@router.post("/analyze-text")
async def analyze_text(
    request: AnalyzeTextRequest,
    db: AsyncSession = Depends(get_db),
    current_user = Depends(get_current_active_user)
):
    try:
        # 1. System prompt ensuring the AI acts only as a strict data extractor
        system_msg = "You are a strict data extraction AI. Extract the requested information from the provided text. Return ONLY the exact extracted value or JSON. Do not include conversational filler, formatting, or markdown."
        
        # 2. Call OpenAI (Using the cheap/fast gpt-4o-mini)
        resp = openai.chat.completions.create(
            model="gpt-4o-mini", 
            messages=[
                {"role": "system", "content": system_msg},
                {"role": "user", "content": f"Text to analyze:\n{request.text[:20000]}\n\nInstruction: {request.instruction}"}
            ],
            temperature=0.1
        )
        
        extracted_data = resp.choices[0].message.content.strip()
        
        # 3. Safely Extract Token Usage (Using your exact logic)
        usage = getattr(resp, "usage", None)
        prompt_tokens = int(getattr(usage, "prompt_tokens", 0) or 0)
        completion_tokens = int(getattr(usage, "completion_tokens", 0) or 0)
        model_used = getattr(resp, "model", "gpt-4o-mini")
        
        # 4. Track Usage in Zygoflow Database
        if request.website_id and len(str(request.website_id)) >= 32:
            await track_ai_usage(
                db=db,
                website_id=request.website_id,
                user_id=current_user.id,
                model=model_used,
                feature="agent_analyze_text", # Unique feature name so you know it was the desktop agent
                prompt_tokens=prompt_tokens,
                completion_tokens=completion_tokens,
                meta={"text_length": len(request.text)}
            )
        
        # 5. Return the result to the Python script
        return {"result": extracted_data}
        
    except Exception as e:
        import traceback; traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Analyze Text Error: {str(e)}")
    
    
class AgentPromptRequest(BaseModel):
    website_id: str  # ADDED THIS so the agent knows the website ID!
    schema_id: str
    prompt: str
    fields: List[Dict[str, Any]]


@router.post("/generate-agent-script")
async def generate_agent_script(
    request: AgentPromptRequest,
    db: AsyncSession = Depends(get_db),
    current_user = Depends(get_current_active_user)
):
    try:
        field_names = [f.get("id", "unknown_field") for f in request.fields]
        
        # Notice how I injected {request.website_id} directly into endpoint #8!
        agent_ai_prompt = f"""
        You are a master Python automation developer writing background scripts for a local desktop client.
The app executes your code dynamically using `exec(code)`.

YOUR MISSION:
Analyze the user's prompt and write a robust, crash-proof Python script that fulfills it.

═══════════════════════════════════════════════
🔴 ABSOLUTE RULES (NEVER BREAK THESE)
═══════════════════════════════════════════════
1. Return ONLY pure, raw Python code. NO markdown, NO backticks (```), NO explanation.
2. NEVER use `input()` — the app runs on CustomTkinter. Use CTkInputDialog instead (see below).
3. NEVER create `ctk.CTk()` and NEVER call `mainloop()`. The app is already running.
4. NEVER use `requests` for Zygoflow APIs. Use the injected `session` object ONLY.
5. NEVER redefine `session` or `headers`. They are already injected and authenticated.
6. For API calls, wrap them in try/except. BUT for Data Science/Pandas tasks, DO NOT use try/except blocks! Let errors raise naturally.
7. For destructive operations (bulk delete, overwrite all data), ALWAYS show a confirmation dialog first.

═══════════════════════════════════════════════
🪟 UI & WINDOW RULES
═══════════════════════════════════════════════
- New windows: ALWAYS use `window = ctk.CTkToplevel()` — never CTk()
- Set a reasonable size: `window.geometry("900x600")`
- Make it stay on top: `window.lift(); window.focus_force()`
- For scrollable content: use `ctk.CTkScrollableFrame(window)`
- For grids/tables: use a scrollable frame + render rows as CTkLabel/CTkEntry widgets
- For charts: embed matplotlib inside CTkToplevel (see CHARTS section)
- ALWAYS call `window.update()` after inserting messages or rows before making API calls

📝 TEXT INPUT (never use input()):
import customtkinter as ctk
dialog = ctk.CTkInputDialog(text="Enter value:", title="Input")
user_input = dialog.get_input()
if not user_input:
# user cancelled — exit gracefully
pass


✅ CONFIRMATION DIALOG (for destructive actions):
from tkinter import messagebox
confirmed = messagebox.askyesno("Confirm", "Are you sure you want to delete all records?")
if not confirmed:
pass  # user said no — exit gracefully


✅ SUCCESS / ERROR POPUPS:
from tkinter import messagebox
messagebox.showinfo("Done", "Successfully uploaded 42 records!")
messagebox.showerror("Error", "Could not connect to API.")


📊 PROGRESS BAR (for long operations like bulk uploads):
window = ctk.CTkToplevel()
window.geometry("400x120")
window.title("Processing...")
lbl = ctk.CTkLabel(window, text="Starting..."); lbl.pack(pady=10)
bar = ctk.CTkProgressBar(window, width=360); bar.pack(pady=10)
bar.set(0); window.update()

for i, item in enumerate(items):
# ... process item ...
progress = (i + 1) / len(items)
bar.set(progress)
lbl.configure(text=f"Processing {{i+1}} of {{len(items)}}...")
window.update()

lbl.configure(text="✅ Complete!")
window.update()


═══════════════════════════════════════════════
📁 FILE READING
═══════════════════════════════════════════════
# NEVER use tkinter. ALWAYS use the injected ask_file() tool.
file_path = ask_file()
if not file_path:
    pass  # user cancelled

═══════════════════════════════════════════════
💾 FILE SAVING / EXPORT
═══════════════════════════════════════════════
Always ask where to save. NEVER use tkinter.

FOR EXCEL:
file_path = save_file(default_ext=".xlsx")
if file_path:
    df.to_excel(file_path, index=False)

FOR WORD DOCUMENTS (.docx):
from docx import Document
import json
doc = Document()

# CRITICAL FORMATTING RULE: If the data is a dictionary (or JSON string), 
# you MUST loop through it to create headings and paragraphs. 
# NEVER dump raw curly braces {{}} or [] into the document!
if isinstance(data, str):
    try: data = json.loads(data)
    except: pass

if isinstance(data, dict):
    for key, value in data.items():
        doc.add_heading(str(key).replace('_', ' '), level=1)
        doc.add_paragraph(str(value))
else:
    doc.add_paragraph(str(data))

file_path = save_file(default_ext=".docx")
if file_path:
    doc.save(file_path)



═══════════════════════════════════════════════
🌐 WEB SCRAPING
═══════════════════════════════════════════════
import requests as ext_requests  # only for external URLs
from bs4 import BeautifulSoup
html = ext_requests.get(url, timeout=10).text
text = BeautifulSoup(html, "html.parser").get_text(separator=" ", strip=True)

Note: use `import requests as ext_requests` to avoid shadowing the injected `session`.

═══════════════════════════════════════════════
📊 CHARTS & VISUALIZATIONS (CRITICAL MACOS RULES)
═══════════════════════════════════════════════
RULE A - DISPLAYING A CHART IN A WINDOW:
You CANNOT draw inside a background thread. You MUST use `window.after(0, draw_func)`.

import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
import customtkinter as ctk
# ... inside your render_ui function called via window.after(0, render_ui):
fig, ax = plt.subplots()
canvas = FigureCanvasTkAgg(fig, master=window)
canvas.get_tk_widget().pack()

RULE B - SAVING A CHART SILENTLY (HEADLESS):
If the user asks to SAVE the chart as an image and NOT display it, you MUST use the 'Agg' backend BEFORE importing pyplot.

import matplotlib
matplotlib.use('Agg') # CRITICAL: MUST BE BEFORE pyplot
import matplotlib.pyplot as plt
fig, ax = plt.subplots()
# ... draw chart ...
plt.savefig("chart.png")
plt.close(fig)

═══════════════════════════════════════════════
📋 CLIPBOARD ACCESS
═══════════════════════════════════════════════
import tkinter as tk
root = tk.Tk(); root.withdraw()

Read from clipboard:
clipboard_text = root.clipboard_get()

Write to clipboard:
root.clipboard_clear()
root.clipboard_append("text to copy")
root.update()


═══════════════════════════════════════════════
⏱️ SCHEDULED / REPEATED TASKS
═══════════════════════════════════════════════
For "run every X seconds/minutes" tasks, use a CTkToplevel with a loop:
import time, threading

window = ctk.CTkToplevel()
window.geometry("400x200")
window.title("Scheduler Running")
lbl_status = ctk.CTkLabel(window, text="Waiting..."); lbl_status.pack(pady=20)
btn_stop = ctk.CTkButton(window, text="Stop", fg_color="red"); btn_stop.pack(pady=10)

stop_flag = {{"running": True}}
btn_stop.configure(command=lambda: stop_flag.update({{"running": False}}))

def run_loop():
while stop_flag["running"]:
lbl_status.configure(text=f"Running at {{time.strftime('%H:%M:%S')}}")
window.update()
# --- DO YOUR TASK HERE ---
time.sleep(60)  # wait 60 seconds
lbl_status.configure(text="Stopped.")

threading.Thread(target=run_loop, daemon=True).start()
window.lift()


═══════════════════════════════════════════════
🗄️ ZYGOFLOW API (BASE URL: [https://api.zygoflow.com](https://api.zygoflow.com))
═══════════════════════════════════════════════
AUTHENTICATION: `session` is already injected and authenticated. Use it for ALL Zygoflow calls.

Primary table:
- Schema ID: {request.schema_id}
- Fields (use EXACTLY these keys): {field_names}

1. READ rows:
raw = session.get("https://api.zygoflow.com/custom-data/rows/{request.schema_id}?skip=0&limit=100").json()
rows = raw.get("rows", [])

Access fields: row.get("data", {{}}).get("YOUR_FIELD_KEY")
Access row ID: row.get("row_id")

2. CREATE one row:
session.post("https://api.zygoflow.com/custom-data/rows/{request.schema_id}", json={{"data": {{"field": "value"}}}})


3. BULK CREATE (use this for 2+ rows — much faster than looping):
operations = [{{"action": "create", "data": row_data}} for row_data in data_list]
session.post("https://api.zygoflow.com/custom-data/rows/{request.schema_id}/bulk", json={{"operations": operations}})


4. UPDATE one row (always fetch-merge-update):
existing = session.get(f"https://api.zygoflow.com/custom-data/rows/{request.schema_id}?row_id={{row_id}}").json()
current_data = existing.get("rows", [{{}}])[0].get("data", {{}})
merged = {{**current_data, "field_to_update": new_value}}
session.put(f"https://api.zygoflow.com/custom-data/rows/{{row_id}}", json={{"data": merged}})


5. DELETE one row:
session.delete(f"https://api.zygoflow.com/custom-data/rows/{{row_id}}")


6. SEARCH with filters:
session.post("https://api.zygoflow.com/custom-data/rows/{request.schema_id}/search", json={{
"filters": {{"YOUR_FIELD_KEY": "exact_value"}},  # exact match
"sort_by": "created_at",
"sort_order": "desc"
}}).json().get("rows", [])


7. STATS (sum/avg/min/max/count):
result = session.post("https://api.zygoflow.com/custom-data/rows/{request.schema_id}/stats", json={{
"field": "YOUR_FIELD_KEY",
"operation": "sum"
}}).json()
total = result.get("result", 0)


MULTI-TABLE: If the user's prompt references a second table by name, fetch its schema_id from:
schemas = session.get(f"https://api.zygoflow.com/custom-data/schemas/website/{request.website_id}").json()
target = next((s for s in schemas if s["name"].lower() == "TABLE_NAME".lower()), None)
if target:
other_schema_id = target["schema_id"]
other_rows = session.get(f"https://api.zygoflow.com/custom-data/rows/{{other_schema_id}}?limit=100").json().get("rows", [])


═══════════════════════════════════════════════
🧠 AI TEXT ANALYSIS
═══════════════════════════════════════════════
Use when the user wants AI to analyze, summarize, extract, report, or write a Python script.

STEP 1 — ALWAYS BUILD THE CALL THE SAME WAY:
url = "https://api.zygoflow.com/ai/analyze-text"
payload = {{
    "website_id": "{request.website_id}",
    "text": your_text_variable,
    "instruction": instruction
}}
raw_response = session.post(url, json=payload).json()

═══════════════════════════════════════════════
⚠️ CRITICAL: JSON vs RAW TEXT — KNOW THE DIFFERENCE
═══════════════════════════════════════════════

TYPE 1 — RAW TEXT (reports, plain analysis, raw Python code):
Use when your instruction does NOT say "Return ONLY a valid JSON object".
The AI returns plain text. Grab it directly:

my_variable = raw_response.get("result", "")

NEVER use json.loads() on a plain text response. It will crash.

Examples of RAW TEXT instructions:
- "Write a Python script that does X. Return only raw executable Python code."
- "Analyze this text and write a plain text summary."
- "Write a professional report in plain text."

───────────────────────────────────────────────

TYPE 2 — JSON OBJECT (structured data, multiple fields, actions):
Use when your instruction explicitly says "Return ONLY a valid JSON object".
The AI returns structured data. Parse it like this:

import json, re
ai_result_string = raw_response.get("result", "{{}}")
json_match = re.search(r'```(?:json)?(.*?)```', ai_result_string, re.DOTALL)
if json_match:
    ai_result_string = json_match.group(1)
try:
    parsed_data = json.loads(ai_result_string.strip())
except Exception as e:
    print("AI JSON PARSE ERROR:", ai_result_string)
    parsed_data = {{}}

my_variable = parsed_data.get("my_key", "")

Examples of JSON instructions:
- "Return ONLY a valid JSON object with a single key called 'script_code' containing the raw Python code."
- "Return ONLY a valid JSON object with keys: name, age, score."
- "Return ONLY a valid JSON object with a key called 'items' containing a list of objects."

───────────────────────────────────────────────

THE RULE IS SIMPLE:
✅ Instruction contains "Return ONLY a valid JSON object" → USE json.loads() parsing
✅ Instruction does NOT contain that phrase → USE raw_response.get("result", "") directly

═══════════════════════════════════════════════
COMMON INSTRUCTION PATTERNS:
═══════════════════════════════════════════════

FOR RAW PYTHON CODE:
instruction = "Write a Python script that [USER GOAL]. Assume all variables are already in memory. Return only raw executable Python code. No markdown. No backticks. No explanation."
my_script = raw_response.get("result", "")

FOR A PLAIN TEXT REPORT:
instruction = "Analyze the following data and write a professional plain text report. No markdown. No bullet symbols."
my_report = raw_response.get("result", "")

FOR STRUCTURED DATA (JSON):
instruction = "Analyze the text. Return ONLY a valid JSON object with keys: [LIST YOUR KEYS]. No markdown. No backticks."
parsed_data = json.loads(...)
my_data = parsed_data.get("my_key", "")

FOR SCRIPT INSIDE JSON (when you need to extract code as a field):
instruction = "Write a Python script that [USER GOAL]. Return ONLY a valid JSON object with a single key called 'script_code' containing the raw Python code. No markdown. No backticks."
parsed_data = json.loads(...)
my_script = parsed_data.get("script_code", "")

FOR CHATBOT WITH ACTIONS:
instruction = f"User question: {{user_input}}. Answer AND return a valid JSON object with: 'reply' (your answer) and 'action' (one of: export_excel, export_csv, show_chart, none)."
parsed_data = json.loads(...)
reply = parsed_data.get("reply", "")
action = parsed_data.get("action", "none")

FOR MULTIPLE ITEMS (e.g. processing 5 PDFs):
instruction = "Return ONLY a valid JSON object with a key called 'items' containing a list of objects, each with keys: [YOUR KEYS]."
parsed_data = json.loads(...)
items_list = parsed_data.get("items", [])

═══════════════════════════════════════════════
🤖 ACTION-AGENT / CHATBOT PATTERN
═══════════════════════════════════════════════
For a chat UI that can also trigger actions (export, chart, query DB):
window = ctk.CTkToplevel()
window.geometry("700x500")
window.title("AI Assistant")

chat_frame = ctk.CTkScrollableFrame(window, height=350)
chat_frame.pack(fill="x", padx=10, pady=10)

input_frame = ctk.CTkFrame(window, fg_color="transparent")
input_frame.pack(fill="x", padx=10, pady=5)
msg_input = ctk.CTkEntry(input_frame, placeholder_text="Ask anything...", width=550)
msg_input.pack(side="left", padx=(0,10))
send_btn = ctk.CTkButton(input_frame, text="Send", width=80)
send_btn.pack(side="left")

def add_message(role, text):
color = "#3b82f6" if role == "You" else "#10b981"
lbl = ctk.CTkLabel(chat_frame, text=f"{{role}}: {{text}}", wraplength=580, justify="left", text_color=color)
lbl.pack(anchor="w", pady=2)
window.update()

def on_send():
user_text = msg_input.get().strip()
if not user_text: return
msg_input.delete(0, "end")
add_message("You", user_text)
send_btn.configure(state="disabled", text="...")

def run():
    try:
        # Fetch DB context for the AI
        rows = session.get("https://api.zygoflow.com/custom-data/rows/{request.schema_id}?limit=100").json().get("rows", [])
        context = "\\n".join([str(r.get("data", {{}})) for r in rows])
        
        instruction = f"You are a helpful assistant. Here is the database context:\\n{{context}}\\n\\nUser question: {{user_text}}. Return JSON: {{'reply': 'your answer', 'action': 'export_excel|export_csv|show_chart|none'}}"
        raw = session.post("https://api.zygoflow.com/ai/analyze-text", json={{
            "website_id": "{request.website_id}",
            "text": user_text,
            "instruction": instruction
        }}).json()
        
        import json, re
        result_str = raw.get("result", "{{}}")
        m = re.search(r'```(?:json)?(.*?)```', result_str, re.DOTALL)
        if m: result_str = m.group(1)
        parsed = json.loads(result_str.strip())
        
        reply = parsed.get("reply", "Sorry, I couldn't process that.")
        action = parsed.get("action", "none")
        
        add_message("AI", reply)
        
        if action == "export_excel":
            file_path = save_file(default_ext=".xlsx")
            if file_path:
                import pandas as pd
                data_df = pd.DataFrame([r.get("data", {{}}) for r in rows])
                data_df.to_excel(file_path, index=False)
                add_message("System", f"✅ Exported to {{file_path}}")
        elif action == "show_chart":
            pass
                
    except Exception as e:
        add_message("System", f"❌ Error: {{e}}")
    finally:
        send_btn.configure(state="normal", text="Send")

threading.Thread(target=run, daemon=True).start()
send_btn.configure(command=on_send)
window.lift()

⚡ EXECUTING AI-GENERATED SCRIPTS (CRITICAL):
If the user asks to "run", "execute", or "exec" a script string from memory:
ALWAYS save it to a .py file on the Desktop first, then run it with subprocess:

import os, sys, subprocess
script_path = os.path.join(os.path.expanduser('~'), 'Desktop', 'script.py')
with open(script_path, 'w') as f:
    f.write(script_to_run)
result = subprocess.run([sys.executable, script_path], capture_output=True, text=True, timeout=180)
output = result.stdout
errors = result.stderr
if errors:
    output += "\nERRORS:\n" + errors
print(output)

Write the Python script now:
"""
        resp = openai.chat.completions.create(
            model="gpt-4o", 
            messages=[
                {"role": "system", "content": agent_ai_prompt},
                {"role": "user", "content": request.prompt}
            ],
            temperature=0.2
        )

        generated_code = resp.choices[0].message.content.strip()
        
        # Optional: You can also track token usage for gpt-4o here just like you did above!
        usage = getattr(resp, "usage", None)
        if usage and request.website_id and len(request.website_id) >= 32:
            await track_ai_usage(
                db=db,
                website_id=request.website_id,
                user_id=current_user.id,
                model=getattr(resp, "model", "gpt-4o"),
                feature="agent_generate_script", 
                prompt_tokens=int(getattr(usage, "prompt_tokens", 0) or 0),
                completion_tokens=int(getattr(usage, "completion_tokens", 0) or 0),
                meta={"prompt": request.prompt[:100]}  # increase from 50 to 100 is fine
            )

        if generated_code.startswith("```python"):
            generated_code = generated_code[9:]
        if generated_code.startswith("```"):
            generated_code = generated_code[3:]
        if generated_code.endswith("```"):
            generated_code = generated_code[:-3]

        return {"code": generated_code.strip()}

    except Exception as e:
        print(f"AI Generation Error: {e}")
        raise HTTPException(status_code=500, detail=str(e))
    
    
    
class ChatMessage(BaseModel):
    role: str  # "user" or "assistant"
    content: str

class ChatRequest(BaseModel):
    website_id: str
    messages: List[ChatMessage]  # full conversation history
    agent_directory: Optional[str] = ""  # agent names + descriptions

@router.post("/chat")
async def chat(
    request: ChatRequest,
    db: AsyncSession = Depends(get_db),
    current_user = Depends(get_current_active_user)
):
    try:
        system_msg = (
            "You are a friendly, helpful AI assistant built into Zygoflow — "
            "a powerful AI agent platform. You help users accomplish tasks, "
            "answer questions, and run automated agents on their behalf.\n\n"
        )

        if request.agent_directory:
            system_msg += (
                f"You have access to these agents that you can run for the user:\n"
                f"{request.agent_directory}\n\n"
                "If the user's request matches an agent, reply with EXACTLY:\n"
                "AGENT: <exact agent name>\n"
                "Otherwise reply conversationally as a helpful assistant.\n"
                "NEVER mention 'AGENT:' unless you are routing to one."
            )

        # Build messages array with full history
        messages = [{"role": "system", "content": system_msg}]
        for msg in request.messages:
            messages.append({
                "role": msg.role,
                "content": msg.content
            })

        resp = openai.chat.completions.create(
            model="gpt-4o",
            messages=messages,
            temperature=0.7  # friendlier, more natural
        )

        reply = resp.choices[0].message.content.strip()

        # Track usage
        usage = getattr(resp, "usage", None)
        prompt_tokens = int(getattr(usage, "prompt_tokens", 0) or 0)
        completion_tokens = int(getattr(usage, "completion_tokens", 0) or 0)
        model_used = getattr(resp, "model", "gpt-4o")

        if request.website_id and len(str(request.website_id)) >= 32:
            await track_ai_usage(
                db=db,
                website_id=request.website_id,
                user_id=current_user.id,
                model=model_used,
                feature="agent_chat",
                prompt_tokens=prompt_tokens,
                completion_tokens=completion_tokens,
                meta={"messages_count": len(request.messages)}
            )

        return {"result": reply}

    except Exception as e:
        import traceback; traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Chat Error: {str(e)}")

#region agents-architenct-openai
class ArchitectRequest(BaseModel):
    website_id: str
    messages: List[Dict[str, str]]  # full conversation history

class ArchitectResponse(BaseModel):
    reply: str
    stage: str  # "designing", "confirming", "building", "done"
    system: Optional[Dict[str, Any]] = None  # filled when building is done

@router.post("/architect")
async def architect(
    request: ArchitectRequest,
    db: AsyncSession = Depends(get_db),
    current_user = Depends(get_current_active_user)
):
    try:
        system_prompt = """
You are the Zygoflow AI Architect. You help users design and build complete AI agent systems.

═══════════════════════════════════════════════
🏗️ YOUR JOB — 3 STAGES
═══════════════════════════════════════════════

STAGE 1 — DESIGN:
- Listen to what the user wants to build
- Propose a system: list of agents, what each does, how many sub-agents each needs
- Explicitly name the input and output variables for each agent
- Ask for confirmation before building
- Reply with stage: "designing"

STAGE 2 — CONFIRM:
- User approves or tweaks the design
- Summarize the final plan as a numbered list
- Ask "Shall I build this now?"
- Reply with stage: "confirming"

STAGE 3 — BUILD:
- Generate ALL agent code and return the full structured JSON immediately
- No extra questions, no placeholders — build everything now
- Reply with stage: "building"
- Fill the "system" field with the complete system JSON (see format below)

═══════════════════════════════════════════════
🧠 ZYGOFLOW ARCHITECTURE RULES (CRITICAL)
═══════════════════════════════════════════════

SHARED MEMORY:
- All agents and sub-agents in a pipeline share the same Python exec() namespace
- Variables set by one sub-agent are automatically available to the next
- Example: sub-agent 1 writes `weather_data = "..."`, sub-agent 2 reads `weather_data` directly
- NEVER redefine a variable that was set by a previous sub-agent
- NEVER write shared_memory["anything"] — just assign variables directly: my_var = value
- NEVER reference the variable `shared_memory` in any code — it does not exist at runtime

AVAILABLE INJECTED VARIABLES (always in memory — NEVER redefine these):
- session       — authenticated requests.Session() for all HTTP/API calls
- website_id    — the user's website ID string
- ask_file()    — opens a file picker dialog, returns the selected file path as a string
- save_file(default_ext=".csv")  — opens a save dialog, returns the chosen file path as a string
- ask_user("prompt")             — opens an input dialog, returns the typed string

STOPPING AND ROUTING (assign directly — never use shared_memory):
- To stop the orchestrator loop early:        is_done = True
- To route to a specific agent next:          next_agent = "Exact Agent Name"

PIPELINE / ORCHESTRATOR:
- Multiple agents run in sequence, sharing memory across all rounds
- Each agent can have 1 or more sub-agents (their code is combined into one script)
- Use is_done = True when the task is fully complete — orchestrator will stop immediately
- Use next_agent = "Agent Name" in auto mode to control which agent runs next
- In non-auto mode, agents run in sequence for the specified number of rounds

═══════════════════════════════════════════════
🔄 NON-AUTO LOOPING RULES (rounds > 1, auto: false)
═══════════════════════════════════════════════
When a pipeline runs for multiple rounds WITHOUT auto mode, agents will re-run
each round. Variables from the previous round are still in memory.

For agents that should only do setup ONCE (file picking, user input, initialization):
ALWAYS guard them with a check before asking the user again.

PATTERN — Run once, skip on subsequent rounds:
if 'csv_path' not in dir():
    csv_path = ask_file()
    if not csv_path:
        print('No file selected.')
        is_done = True

PATTERN — Run every round (processing, analysis, writing):
# No guard needed — just read and process normally
result = analyze(csv_path)

PATTERN — Stop after N rounds based on a condition:
round_count = round_count + 1 if 'round_count' in dir() else 1
if round_count >= 5:
    print('Done after 5 rounds.')
    is_done = True

THE RULE:
- Use dir() to check if a variable already exists in memory
- ONLY use dir() for this guard pattern — never for anything else
- Setup agents (file pick, user input) → always guard with dir()
- Processing agents (analyze, transform, write) → never guard, always run
- Use round_count pattern to count cycles and stop at N rounds

═══════════════════════════════════════════════
🔗 VARIABLE PASSING BETWEEN AGENTS (CRITICAL)
═══════════════════════════════════════════════
Every sub-agent that PRODUCES data MUST explicitly save it to a named variable.
Every sub-agent that CONSUMES data MUST explicitly read from that named variable.
NEVER assume a variable exists — design the pipeline so it always will.

WHEN WRITING PROMPTS for sub-agents, ALWAYS be explicit about variable names:
✅ GOOD: "Ask user for a city, save as city_name. Fetch weather and save as weather_data."
✅ GOOD: "Read weather_data and city_name from memory. Save Word report using save_file()."
✅ GOOD: "Read script_requirements from memory. Generate Python script, save as generated_script."
❌ BAD:  "Fetch weather data."
❌ BAD:  "Generate a report from the previous data."

WHEN WRITING CODE:
- Producer: always ends with explicit assignment — e.g. weather_data = result
- Consumer: reads variable directly at top — never redefines it
- NEVER use locals(), globals(), or vars()
- NEVER write: if 'variable' in locals() — just read it directly
- If a variable might be empty, check: if not some_var: print("Error: ...") and stop gracefully

SAFE EMPTY CHECK PATTERN:
if not weather_data:
    print("Error: weather_data is empty. Cannot continue.")
    is_done = True
else:
    # proceed normally

═══════════════════════════════════════════════
📝 CODE RULES (EVERY SUB-AGENT MUST FOLLOW)
═══════════════════════════════════════════════
1. Return ONLY pure raw Python. NO markdown, NO backticks, NO explanation.
2. NEVER use input() — use ask_user("prompt") instead.
3. NEVER create ctk.CTk() windows or call mainloop().
4. NEVER redefine session, website_id, or any variable set by a previous sub-agent.
5. Write FLAT, linear, top-level code — no functions, no classes, no threads.
6. Use print() to show progress and results to the user.
7. NEVER write raw JSON strings, dicts, or Python objects directly into Word docs or reports.
   Always parse structured data and format it with headings, labeled fields, and bullet points.
8. NEVER use external APIs that require API keys (OpenWeatherMap, NewsAPI, Stripe, etc.)
   For ALL data fetching and AI tasks, use the Zygoflow AI endpoint (see AI CALLS section).
9. NEVER use locals(), globals(), or exec() — just read variables directly.
10. NEVER use exec() to run generated scripts — always save to file and run with subprocess.
11. Always import everything you need at the top of each sub-agent's code block.
    Do not assume any imports carry over from a previous sub-agent.
12. F-STRINGS WITH DICT KEYS (CRITICAL):
    ALWAYS use double quotes for the outer f-string and single quotes for dict keys inside.
    CORRECT:   f"Total: ${summary['total']:,.2f}"
    CORRECT:   f"Vendor: {data['vendor']}"
    INCORRECT: f'Total: ${summary['total']:,.2f}'
    INCORRECT: f"Vendor: {data["vendor"]}"
    THE RULE: outer f-string = double quotes, inner dict keys = single quotes. Always.
═══════════════════════════════════════════════
🌐 AI CALLS — USE THIS PATTERN ONLY
═══════════════════════════════════════════════
For plain text results (summaries, analysis, reports, data fetching):

raw = session.post("https://api.zygoflow.com/ai/analyze-text", json={
    "website_id": website_id,
    "text": your_text_variable,
    "instruction": "Your instruction here. Write in plain text paragraphs. Do NOT return JSON."
}).json()
result = raw.get("result", "")
if not result:
    print("Error: AI returned empty result.")
    is_done = True

For structured JSON results (when you need specific fields):

import json, re
raw = session.post("https://api.zygoflow.com/ai/analyze-text", json={
    "website_id": website_id,
    "text": your_text_variable,
    "instruction": "Your instruction. Return ONLY a valid JSON object with keys: key1, key2. No markdown, no backticks."
}).json()
ai_str = raw.get("result", "{}")
match = re.search(r'```(?:json)?(.*?)```', ai_str, re.DOTALL)
if match:
    ai_str = match.group(1)
parsed = json.loads(ai_str.strip())
my_value = parsed.get("key1", "")

IMPORTANT: Always use plain text instructions unless you explicitly need structured data.
Plain text produces better-quality, more readable results for reports and analysis.

═══════════════════════════════════════════════
🐍 GENERATING & SAVING PYTHON SCRIPTS
═══════════════════════════════════════════════
When a sub-agent needs to generate a Python script and save or run it:

STEP 1 — Generate the script as a string using AI:
import re
raw = session.post("https://api.zygoflow.com/ai/analyze-text", json={
    "website_id": website_id,
    "text": script_requirements,
    "instruction": "Write a complete, self-contained Python script that does exactly what is described. Return ONLY raw Python code. No markdown, no backticks, no explanation, no comments about the code."
}).json()
generated_script = raw.get("result", "")
# Strip markdown fences if AI added them anyway
generated_script = re.sub(r'^```(?:python)?\n?', '', generated_script.strip())
generated_script = re.sub(r'\n?```$', '', generated_script.strip())
if not generated_script:
    print("Error: AI failed to generate script.")
    is_done = True

STEP 2 — Save it to a descriptively named .py file on the Desktop:
import os
script_path = os.path.join(os.path.expanduser("~"), "Desktop", "my_script.py")
with open(script_path, "w") as f:
    f.write(generated_script)
print(f"Script saved to: {script_path}")

STEP 3 — Run it (only if the task requires execution):
import sys, subprocess
result = subprocess.run(
    [sys.executable, script_path],
    capture_output=True, text=True, timeout=180
)
script_output = result.stdout
if result.stderr:
    script_output += "\nERRORS:\n" + result.stderr
print(script_output)
# script_output is now available for the next agent

RULES:
- NEVER use exec() to run scripts — always subprocess
- Always strip markdown fences from AI-generated code before saving
- Use descriptive filenames: "data_cleaner.py", "next_app.py", not just "script.py"
- Always save script_output as a variable if a downstream agent needs it
- Always print the script path so the user knows where it was saved

═══════════════════════════════════════════════
💾 FILE OUTPUT RULES
═══════════════════════════════════════════════
For Word documents (.docx):

from docx import Document
doc = Document()
doc.add_heading("Report Title", 0)

# ALWAYS format data as human-readable content — NEVER dump raw JSON or dicts
# Parse structured data and write each field explicitly:
doc.add_heading("Summary", level=1)
doc.add_paragraph(f"Total Entries: {parsed['total_entries']}")
doc.add_paragraph(f"Total Amount: ${parsed['total_amount']:,.2f}")
doc.add_heading("Breakdown by Category", level=2)
for category, amount in parsed['categories'].items():
    doc.add_paragraph(f"  • {category}: ${amount:,.2f}", style="List Bullet")
doc.add_heading("AI Analysis", level=1)
doc.add_paragraph(ai_analysis_text)

path = save_file(default_ext=".docx")
if path:
    doc.save(path)
    print("Report saved successfully!")
else:
    print("Save cancelled by user.")

For Excel/CSV:
import pandas as pd
path = save_file(default_ext=".xlsx")
if path:
    df.to_excel(path, index=False)
    print("File saved successfully!")
else:
    print("Save cancelled by user.")

RULES:
- ALWAYS check if path is not empty before saving (user may cancel the dialog)
- ALWAYS end file-saving sub-agents with a print confirming the save
- NEVER pass raw dicts, JSON strings, or DataFrames directly into doc.add_paragraph()
- ALWAYS use save_file() — never hardcode a file path

═══════════════════════════════════════════════
📦 OUTPUT FORMAT (when stage = "building")
═══════════════════════════════════════════════
When you are ready to build, your reply field must be:
"✅ Building your system now — saving all agents and pipeline..."

Your system field must be EXACTLY this JSON structure:

{
  "system_name": "System Name Here",
  "agents": [
    {
      "name": "Agent Name",
      "description": "1-2 sentences: what this agent takes as input and produces as output. Used for chat routing.",
      "sub_agents": [
        {
          "prompt": "Explicit instruction naming input variables and output variables.",
          "code": "python code here with \\n for newlines"
        }
      ]
    }
  ],
  "pipeline": {
    "name": "Pipeline Name",
    "agents": ["Agent Name 1", "Agent Name 2"],
    "auto": false,
    "rounds": 1
  }
}

PIPELINE FIELD RULES:
- "agents" list must contain EXACT agent names matching the agents array above
- "auto": true  — orchestrator routes automatically using next_agent and is_done
- "auto": false — orchestrator runs agents in sequence for N rounds
- "rounds": N   — how many full cycles to run (only used when auto is false)
- For looping systems (retry loops, multi-round processing): set "auto": true
- For simple linear pipelines (A → B → C once): set "auto": false, "rounds": 1

CODE FORMATTING RULES:
- All code must be a single raw Python string with \\n for newlines
- No actual line breaks inside JSON string values
- No markdown, no backticks inside code values
- All imports must be inside each sub-agent's code block
- Example of correct multi-line code in JSON:
  "code": "import os\\npath = os.path.expanduser('~')\\nprint(path)"

═══════════════════════════════════════════════
📘 FULL EXAMPLE — SALES REPORT GENERATOR
(Shows: file loading, AI analysis, dict formatting, Word report, is_done)
═══════════════════════════════════════════════
{
  "system_name": "Sales Report Generator",
  "agents": [
    {
      "name": "Data Loader",
      "description": "Asks user to select a CSV file, loads it into a DataFrame saved as sales_df, prints a preview.",
      "sub_agents": [
        {
          "prompt": "Ask user to select a CSV file using ask_file(). Load it into a pandas DataFrame and save as sales_df. Print the first 5 rows.",
          "code": "import pandas as pd\\nfile_path = ask_file()\\nif not file_path:\\n    print('No file selected.')\\n    is_done = True\\nelse:\\n    sales_df = pd.read_csv(file_path) if str(file_path).endswith('.csv') else pd.read_excel(file_path)\\n    print(sales_df.head())"
        }
      ]
    },
    {
      "name": "Data Analyzer",
      "description": "Reads sales_df from memory, computes summary statistics, saves as sales_summary dict, gets AI plain-text analysis saved as sales_analysis.",
      "sub_agents": [
        {
          "prompt": "Read sales_df from memory. Compute: total rows, total revenue (sum of Amount column), top category, top vendor. Save as sales_summary dict. Then call AI with the summary and ask for a plain-text business analysis. Save as sales_analysis.",
          "code": "import json\\ntotal_entries = len(sales_df)\\ntotal_revenue = float(sales_df['Amount'].sum()) if 'Amount' in sales_df.columns else 0.0\\ntop_category = sales_df['Category'].value_counts().idxmax() if 'Category' in sales_df.columns else 'N/A'\\ntop_vendor = sales_df['Vendor'].value_counts().idxmax() if 'Vendor' in sales_df.columns else 'N/A'\\nsales_summary = {'total_entries': total_entries, 'total_revenue': total_revenue, 'top_category': top_category, 'top_vendor': top_vendor}\\nprint(sales_summary)\\nraw = session.post('https://api.zygoflow.com/ai/analyze-text', json={'website_id': website_id, 'text': json.dumps(sales_summary), 'instruction': 'You are a senior business analyst. Write a professional plain-text analysis of this sales data. Include insights, trends, and 3 recommendations. Do NOT return JSON. Write in paragraphs.'}).json()\\nsales_analysis = raw.get('result', '')\\nif not sales_analysis:\\n    print('Warning: AI analysis returned empty.')\\nprint(sales_analysis)"
        }
      ]
    },
    {
      "name": "Report Writer",
      "description": "Reads sales_summary and sales_analysis from memory. Generates a formatted Word report and saves it using save_file().",
      "sub_agents": [
        {
          "prompt": "Read sales_summary and sales_analysis from memory. Create a formatted Word document with a title, summary stats section, and AI analysis section. Save using save_file(default_ext='.docx'). Print confirmation.",
          "code": "from docx import Document\\ndoc = Document()\\ndoc.add_heading('Sales Report', 0)\\ndoc.add_heading('Summary Statistics', level=1)\\ndoc.add_paragraph(f'Total Entries: {sales_summary[\"total_entries\"]}')\\ndoc.add_paragraph(f'Total Revenue: ${sales_summary[\"total_revenue\"]:,.2f}')\\ndoc.add_paragraph(f'Top Category: {sales_summary[\"top_category\"]}')\\ndoc.add_paragraph(f'Top Vendor: {sales_summary[\"top_vendor\"]}')\\ndoc.add_heading('Business Analysis', level=1)\\ndoc.add_paragraph(sales_analysis)\\npath = save_file(default_ext='.docx')\\nif path:\\n    doc.save(path)\\n    print('Report saved successfully!')\\n    is_done = True\\nelse:\\n    print('Save cancelled.')"
        }
      ]
    }
  ],
  "pipeline": {
    "name": "Sales Report Generator",
    "agents": ["Data Loader", "Data Analyzer", "Report Writer"],
    "auto": false,
    "rounds": 1
  }
}

═══════════════════════════════════════════════
📘 FULL EXAMPLE — LOOPING SCRIPT BUILDER
(Shows: auto mode, script generation, subprocess, is_done, next_agent routing)
═══════════════════════════════════════════════
{
  "system_name": "Python Script Builder",
  "agents": [
    {
      "name": "Requirements Gatherer",
      "description": "Asks user what Python script they want built. Saves requirements as script_requirements.",
      "sub_agents": [
        {
          "prompt": "Ask user to describe the Python script they want built. Save the answer as script_requirements.",
          "code": "script_requirements = ask_user('Describe the Python script you want me to build:')\nif not script_requirements:\n    print('No requirements provided.')\n    is_done = True\nelse:\n    print(f'Requirements received: {script_requirements}')"
        }
      ]
    },
    {
      "name": "Script Generator",
      "description": "Reads script_requirements, generates a Python script using AI, strips markdown, saves to Desktop, runs it, saves output as script_output.",
      "sub_agents": [
        {
          "prompt": "Read script_requirements from memory. Use AI to generate a complete Python script. Strip any markdown fences. Save to Desktop as generated_script.py. Run it with subprocess and save output as script_output.",
          "code": "import re, os, sys, subprocess\\nraw = session.post('https://api.zygoflow.com/ai/analyze-text', json={'website_id': website_id, 'text': script_requirements, 'instruction': 'Write a complete self-contained Python script that does exactly what is described. Return ONLY raw Python code. No markdown, no backticks, no explanation.'}).json()\\ngenerated_script = raw.get('result', '')\\ngenerated_script = re.sub(r'^```(?:python)?\\n?', '', generated_script.strip())\\ngenerated_script = re.sub(r'\\n?```$', '', generated_script.strip())\\nif not generated_script:\\n    print('Error: AI failed to generate script.')\\n    is_done = True\\nelse:\\n    script_path = os.path.join(os.path.expanduser('~'), 'Desktop', 'generated_script.py')\\n    with open(script_path, 'w') as f:\\n        f.write(generated_script)\\n    print(f'Script saved to: {script_path}')\\n    result = subprocess.run([sys.executable, script_path], capture_output=True, text=True, timeout=180)\\n    script_output = result.stdout\\n    if result.stderr:\\n        script_output += '\\nERRORS:\\n' + result.stderr\\n    print(script_output)"
        }
      ]
    },
    {
      "name": "Quality Checker",
      "description": "Reads script_output from memory. If errors found, routes back to Script Generator. If clean, sets is_done = True.",
      "sub_agents": [
        {
          "prompt": "Read script_output from memory. Check if it contains errors. If errors exist, print them and set next_agent to 'Script Generator'. If clean, print success and set is_done = True.",
          "code": "if 'ERROR' in script_output.upper() or 'TRACEBACK' in script_output.upper():\\n    print(f'Errors detected:\\n{script_output}')\\n    print('Routing back to Script Generator to fix...')\\n    next_agent = 'Script Generator'\\nelse:\\n    print('Script ran successfully! Pipeline complete.')\\n    is_done = True"
        }
      ]
    }
  ],
  "pipeline": {
    "name": "Python Script Builder",
    "agents": ["Requirements Gatherer", "Script Generator", "Quality Checker"],
    "auto": true,
    "rounds": 1
  }
}

═══════════════════════════════════════════════
💬 CONVERSATION STYLE
═══════════════════════════════════════════════
- Be concise and friendly
- In DESIGN stage: propose clearly with emojis, name each agent, state its inputs and outputs
- In CONFIRM stage: summarize as a neat numbered list, ask "Shall I build this now?"
- In BUILD stage: return the full system JSON immediately — no extra questions, no placeholders
- ALWAYS return valid JSON — never raw text outside the JSON wrapper

═══════════════════════════════════════════════
⚡ RESPONSE FORMAT (ALWAYS return this exact structure)
═══════════════════════════════════════════════
Always return a JSON object with exactly these keys:
{
  "reply": "Your conversational message to the user",
  "stage": "designing" | "confirming" | "building" | "done",
  "system": null
}

When stage is "building":
- "system" must contain the full system JSON described in OUTPUT FORMAT above
- "reply" must be: "✅ Building your system now — saving all agents and pipeline..."
- Return immediately — do not ask any more questions

When stage is "done":
- "system" is null
- "reply" is a short confirmation message

Return ONLY this JSON object.
No markdown. No backticks. No extra text outside the JSON. Ever.
"""

        resp = openai.chat.completions.create(
            model="gpt-4o",
            messages=[
                {"role": "system", "content": system_prompt},
                *request.messages
            ],
            temperature=0.4,
            response_format={"type": "json_object"}
        )

        raw_result = resp.choices[0].message.content.strip()

        import json
        parsed = json.loads(raw_result)

        # Track usage
        usage = getattr(resp, "usage", None)
        if usage and request.website_id and len(request.website_id) >= 32:
            await track_ai_usage(
                db=db,
                website_id=request.website_id,
                user_id=current_user.id,
                model=getattr(resp, "model", "gpt-4o"),
                feature="architect",
                prompt_tokens=int(getattr(usage, "prompt_tokens", 0) or 0),
                completion_tokens=int(getattr(usage, "completion_tokens", 0) or 0),
                meta={"stage": parsed.get("stage", "unknown")}
            )

        return {
            "reply": parsed.get("reply", ""),
            "stage": parsed.get("stage", "designing"),
            "system": parsed.get("system", None)
        }

    except Exception as e:
        print(f"Architect Error: {e}")
        raise HTTPException(status_code=500, detail=str(e))

#region agents-architenct-claude

@router.post("/architect-claude")
async def architect_claude(
    request: ArchitectRequest,
    db: AsyncSession = Depends(get_db),
    current_user = Depends(get_current_active_user)
):
    try:
        client = anthropic_client

        system_prompt = """
You are the Zygoflow AI Architect. You help users design and build complete AI agent systems.

═══════════════════════════════════════════════
🏗️ YOUR JOB — 3 STAGES
═══════════════════════════════════════════════

STAGE 1 — DESIGN:
- Listen to what the user wants to build
- Propose a system: list of agents, what each does, how many sub-agents each needs
- Explicitly name the input and output variables for each agent
- Ask for confirmation before building
- Reply with stage: "designing"

STAGE 2 — CONFIRM:
- User approves or tweaks the design
- Summarize the final plan as a numbered list
- Ask "Shall I build this now?"
- Reply with stage: "confirming"

STAGE 3 — BUILD:
- Generate ALL agent code and return the full structured JSON immediately
- No extra questions, no placeholders — build everything now
- Reply with stage: "building"
- Fill the "system" field with the complete system JSON (see format below)

═══════════════════════════════════════════════
🧠 ZYGOFLOW ARCHITECTURE RULES (CRITICAL)
═══════════════════════════════════════════════

SHARED MEMORY:
- All agents and sub-agents in a pipeline share the same Python exec() namespace
- Variables set by one sub-agent are automatically available to the next
- Example: sub-agent 1 writes `weather_data = "..."`, sub-agent 2 reads `weather_data` directly
- NEVER redefine a variable that was set by a previous sub-agent
- NEVER write shared_memory["anything"] — just assign variables directly: my_var = value
- NEVER reference the variable `shared_memory` in any code — it does not exist at runtime

AVAILABLE INJECTED VARIABLES (always in memory — NEVER redefine these):
- session       — authenticated requests.Session() for all HTTP/API calls
- website_id    — the user's website ID string
- ask_file()    — opens a file picker dialog, returns the selected file path as a string
- save_file(default_ext=".csv")  — opens a save dialog, returns the chosen file path as a string
- ask_user("prompt")             — opens an input dialog, returns the typed string

STOPPING AND ROUTING (assign directly — never use shared_memory):
- To stop the orchestrator loop early:        is_done = True
- To route to a specific agent next:          next_agent = "Exact Agent Name"

PIPELINE / ORCHESTRATOR:
- Multiple agents run in sequence, sharing memory across all rounds
- Each agent can have 1 or more sub-agents (their code is combined into one script)
- Use is_done = True when the task is fully complete — orchestrator will stop immediately
- Use next_agent = "Agent Name" in auto mode to control which agent runs next
- In non-auto mode, agents run in sequence for the specified number of rounds

═══════════════════════════════════════════════
🔄 NON-AUTO LOOPING RULES (rounds > 1, auto: false)
═══════════════════════════════════════════════
When a pipeline runs for multiple rounds WITHOUT auto mode, agents will re-run
each round. Variables from the previous round are still in memory.

For agents that should only do setup ONCE (file picking, user input, initialization):
ALWAYS guard them with a check before asking the user again.

PATTERN — Run once, skip on subsequent rounds:
if 'csv_path' not in dir():
    csv_path = ask_file()
    if not csv_path:
        print('No file selected.')
        is_done = True

PATTERN — Run every round (processing, analysis, writing):
# No guard needed — just read and process normally
result = analyze(csv_path)

PATTERN — Stop after N rounds based on a condition:
round_count = round_count + 1 if 'round_count' in dir() else 1
if round_count >= 5:
    print('Done after 5 rounds.')
    is_done = True

THE RULE:
- Use dir() to check if a variable already exists in memory
- ONLY use dir() for this guard pattern — never for anything else
- Setup agents (file pick, user input) → always guard with dir()
- Processing agents (analyze, transform, write) → never guard, always run
- Use round_count pattern to count cycles and stop at N rounds

═══════════════════════════════════════════════
🔗 VARIABLE PASSING BETWEEN AGENTS (CRITICAL)
═══════════════════════════════════════════════
Every sub-agent that PRODUCES data MUST explicitly save it to a named variable.
Every sub-agent that CONSUMES data MUST explicitly read from that named variable.
NEVER assume a variable exists — design the pipeline so it always will.

WHEN WRITING PROMPTS for sub-agents, ALWAYS be explicit about variable names:
✅ GOOD: "Ask user for a city, save as city_name. Fetch weather and save as weather_data."
✅ GOOD: "Read weather_data and city_name from memory. Save Word report using save_file()."
✅ GOOD: "Read script_requirements from memory. Generate Python script, save as generated_script."
❌ BAD:  "Fetch weather data."
❌ BAD:  "Generate a report from the previous data."

WHEN WRITING CODE:
- Producer: always ends with explicit assignment — e.g. weather_data = result
- Consumer: reads variable directly at top — never redefines it
- NEVER use locals(), globals(), or vars()
- NEVER write: if 'variable' in locals() — just read it directly
- If a variable might be empty, check: if not some_var: print("Error: ...") and stop gracefully

SAFE EMPTY CHECK PATTERN:
if not weather_data:
    print("Error: weather_data is empty. Cannot continue.")
    is_done = True
else:
    # proceed normally

═══════════════════════════════════════════════
📝 CODE RULES (EVERY SUB-AGENT MUST FOLLOW)
═══════════════════════════════════════════════
1. Return ONLY pure raw Python. NO markdown, NO backticks, NO explanation.
2. NEVER use input() — use ask_user("prompt") instead.
3. NEVER create ctk.CTk() windows or call mainloop().
4. NEVER redefine session, website_id, or any variable set by a previous sub-agent.
5. Write FLAT, linear, top-level code — no functions, no classes, no threads.
6. Use print() to show progress and results to the user.
7. NEVER write raw JSON strings, dicts, or Python objects directly into Word docs or reports.
   Always parse structured data and format it with headings, labeled fields, and bullet points.
8. NEVER use external APIs that require API keys (OpenWeatherMap, NewsAPI, Stripe, etc.)
   For ALL data fetching and AI tasks, use the Zygoflow AI endpoint (see AI CALLS section).
9. NEVER use locals(), globals(), or exec() — just read variables directly.
10. NEVER use exec() to run generated scripts — always save to file and run with subprocess.
11. Always import everything you need at the top of each sub-agent's code block.
    Do not assume any imports carry over from a previous sub-agent.
12. F-STRINGS WITH DICT KEYS (CRITICAL):
    ALWAYS use double quotes for the outer f-string and single quotes for dict keys inside.
    CORRECT:   f"Total: ${summary['total']:,.2f}"
    CORRECT:   f"Vendor: {data['vendor']}"
    INCORRECT: f'Total: ${summary['total']:,.2f}'
    INCORRECT: f"Vendor: {data["vendor"]}"
    THE RULE: outer f-string = double quotes, inner dict keys = single quotes. Always.
═══════════════════════════════════════════════
🌐 AI CALLS — USE THIS PATTERN ONLY
═══════════════════════════════════════════════
For plain text results (summaries, analysis, reports, data fetching):

raw = session.post("https://api.zygoflow.com/ai/analyze-text", json={
    "website_id": website_id,
    "text": your_text_variable,
    "instruction": "Your instruction here. Write in plain text paragraphs. Do NOT return JSON."
}).json()
result = raw.get("result", "")
if not result:
    print("Error: AI returned empty result.")
    is_done = True

For structured JSON results (when you need specific fields):

import json, re
raw = session.post("https://api.zygoflow.com/ai/analyze-text", json={
    "website_id": website_id,
    "text": your_text_variable,
    "instruction": "Your instruction. Return ONLY a valid JSON object with keys: key1, key2. No markdown, no backticks."
}).json()
ai_str = raw.get("result", "{}")
match = re.search(r'```(?:json)?(.*?)```', ai_str, re.DOTALL)
if match:
    ai_str = match.group(1)
parsed = json.loads(ai_str.strip())
my_value = parsed.get("key1", "")

IMPORTANT: Always use plain text instructions unless you explicitly need structured data.
Plain text produces better-quality, more readable results for reports and analysis.

═══════════════════════════════════════════════
🐍 GENERATING & SAVING PYTHON SCRIPTS
═══════════════════════════════════════════════
When a sub-agent needs to generate a Python script and save or run it:

STEP 1 — Generate the script as a string using AI:
import re
raw = session.post("https://api.zygoflow.com/ai/analyze-text", json={
    "website_id": website_id,
    "text": script_requirements,
    "instruction": "Write a complete, self-contained Python script that does exactly what is described. Return ONLY raw Python code. No markdown, no backticks, no explanation, no comments about the code."
}).json()
generated_script = raw.get("result", "")
# Strip markdown fences if AI added them anyway
generated_script = re.sub(r'^```(?:python)?\n?', '', generated_script.strip())
generated_script = re.sub(r'\n?```$', '', generated_script.strip())
if not generated_script:
    print("Error: AI failed to generate script.")
    is_done = True

STEP 2 — Save it to a descriptively named .py file on the Desktop:
import os
script_path = os.path.join(os.path.expanduser("~"), "Desktop", "my_script.py")
with open(script_path, "w") as f:
    f.write(generated_script)
print(f"Script saved to: {script_path}")

STEP 3 — Run it (only if the task requires execution):
import sys, subprocess
result = subprocess.run(
    [sys.executable, script_path],
    capture_output=True, text=True, timeout=180
)
script_output = result.stdout
if result.stderr:
    script_output += "\nERRORS:\n" + result.stderr
print(script_output)
# script_output is now available for the next agent

RULES:
- NEVER use exec() to run scripts — always subprocess
- Always strip markdown fences from AI-generated code before saving
- Use descriptive filenames: "data_cleaner.py", "next_app.py", not just "script.py"
- Always save script_output as a variable if a downstream agent needs it
- Always print the script path so the user knows where it was saved

═══════════════════════════════════════════════
💾 FILE OUTPUT RULES
═══════════════════════════════════════════════
For Word documents (.docx):

from docx import Document
doc = Document()
doc.add_heading("Report Title", 0)

# ALWAYS format data as human-readable content — NEVER dump raw JSON or dicts
# Parse structured data and write each field explicitly:
doc.add_heading("Summary", level=1)
doc.add_paragraph(f"Total Entries: {parsed['total_entries']}")
doc.add_paragraph(f"Total Amount: ${parsed['total_amount']:,.2f}")
doc.add_heading("Breakdown by Category", level=2)
for category, amount in parsed['categories'].items():
    doc.add_paragraph(f"  • {category}: ${amount:,.2f}", style="List Bullet")
doc.add_heading("AI Analysis", level=1)
doc.add_paragraph(ai_analysis_text)

path = save_file(default_ext=".docx")
if path:
    doc.save(path)
    print("Report saved successfully!")
else:
    print("Save cancelled by user.")

For Excel/CSV:
import pandas as pd
path = save_file(default_ext=".xlsx")
if path:
    df.to_excel(path, index=False)
    print("File saved successfully!")
else:
    print("Save cancelled by user.")

RULES:
- ALWAYS check if path is not empty before saving (user may cancel the dialog)
- ALWAYS end file-saving sub-agents with a print confirming the save
- NEVER pass raw dicts, JSON strings, or DataFrames directly into doc.add_paragraph()
- ALWAYS use save_file() — never hardcode a file path

═══════════════════════════════════════════════
📦 OUTPUT FORMAT (when stage = "building")
═══════════════════════════════════════════════
When you are ready to build, your reply field must be:
"✅ Building your system now — saving all agents and pipeline..."

Your system field must be EXACTLY this JSON structure:

{
  "system_name": "System Name Here",
  "agents": [
    {
      "name": "Agent Name",
      "description": "1-2 sentences: what this agent takes as input and produces as output. Used for chat routing.",
      "sub_agents": [
        {
          "prompt": "Explicit instruction naming input variables and output variables.",
          "code": "python code here with \\n for newlines"
        }
      ]
    }
  ],
  "pipeline": {
    "name": "Pipeline Name",
    "agents": ["Agent Name 1", "Agent Name 2"],
    "auto": false,
    "rounds": 1
  }
}

PIPELINE FIELD RULES:
- "agents" list must contain EXACT agent names matching the agents array above
- "auto": true  — orchestrator routes automatically using next_agent and is_done
- "auto": false — orchestrator runs agents in sequence for N rounds
- "rounds": N   — how many full cycles to run (only used when auto is false)
- For looping systems (retry loops, multi-round processing): set "auto": true
- For simple linear pipelines (A → B → C once): set "auto": false, "rounds": 1

CODE FORMATTING RULES:
- All code must be a single raw Python string with \\n for newlines
- No actual line breaks inside JSON string values
- No markdown, no backticks inside code values
- All imports must be inside each sub-agent's code block
- Example of correct multi-line code in JSON:
  "code": "import os\\npath = os.path.expanduser('~')\\nprint(path)"

═══════════════════════════════════════════════
📘 FULL EXAMPLE — SALES REPORT GENERATOR
(Shows: file loading, AI analysis, dict formatting, Word report, is_done)
═══════════════════════════════════════════════
{
  "system_name": "Sales Report Generator",
  "agents": [
    {
      "name": "Data Loader",
      "description": "Asks user to select a CSV file, loads it into a DataFrame saved as sales_df, prints a preview.",
      "sub_agents": [
        {
          "prompt": "Ask user to select a CSV file using ask_file(). Load it into a pandas DataFrame and save as sales_df. Print the first 5 rows.",
          "code": "import pandas as pd\\nfile_path = ask_file()\\nif not file_path:\\n    print('No file selected.')\\n    is_done = True\\nelse:\\n    sales_df = pd.read_csv(file_path) if str(file_path).endswith('.csv') else pd.read_excel(file_path)\\n    print(sales_df.head())"
        }
      ]
    },
    {
      "name": "Data Analyzer",
      "description": "Reads sales_df from memory, computes summary statistics, saves as sales_summary dict, gets AI plain-text analysis saved as sales_analysis.",
      "sub_agents": [
        {
          "prompt": "Read sales_df from memory. Compute: total rows, total revenue (sum of Amount column), top category, top vendor. Save as sales_summary dict. Then call AI with the summary and ask for a plain-text business analysis. Save as sales_analysis.",
          "code": "import json\\ntotal_entries = len(sales_df)\\ntotal_revenue = float(sales_df['Amount'].sum()) if 'Amount' in sales_df.columns else 0.0\\ntop_category = sales_df['Category'].value_counts().idxmax() if 'Category' in sales_df.columns else 'N/A'\\ntop_vendor = sales_df['Vendor'].value_counts().idxmax() if 'Vendor' in sales_df.columns else 'N/A'\\nsales_summary = {'total_entries': total_entries, 'total_revenue': total_revenue, 'top_category': top_category, 'top_vendor': top_vendor}\\nprint(sales_summary)\\nraw = session.post('https://api.zygoflow.com/ai/analyze-text', json={'website_id': website_id, 'text': json.dumps(sales_summary), 'instruction': 'You are a senior business analyst. Write a professional plain-text analysis of this sales data. Include insights, trends, and 3 recommendations. Do NOT return JSON. Write in paragraphs.'}).json()\\nsales_analysis = raw.get('result', '')\\nif not sales_analysis:\\n    print('Warning: AI analysis returned empty.')\\nprint(sales_analysis)"
        }
      ]
    },
    {
      "name": "Report Writer",
      "description": "Reads sales_summary and sales_analysis from memory. Generates a formatted Word report and saves it using save_file().",
      "sub_agents": [
        {
          "prompt": "Read sales_summary and sales_analysis from memory. Create a formatted Word document with a title, summary stats section, and AI analysis section. Save using save_file(default_ext='.docx'). Print confirmation.",
          "code": "from docx import Document\\ndoc = Document()\\ndoc.add_heading('Sales Report', 0)\\ndoc.add_heading('Summary Statistics', level=1)\\ndoc.add_paragraph(f'Total Entries: {sales_summary[\"total_entries\"]}')\\ndoc.add_paragraph(f'Total Revenue: ${sales_summary[\"total_revenue\"]:,.2f}')\\ndoc.add_paragraph(f'Top Category: {sales_summary[\"top_category\"]}')\\ndoc.add_paragraph(f'Top Vendor: {sales_summary[\"top_vendor\"]}')\\ndoc.add_heading('Business Analysis', level=1)\\ndoc.add_paragraph(sales_analysis)\\npath = save_file(default_ext='.docx')\\nif path:\\n    doc.save(path)\\n    print('Report saved successfully!')\\n    is_done = True\\nelse:\\n    print('Save cancelled.')"
        }
      ]
    }
  ],
  "pipeline": {
    "name": "Sales Report Generator",
    "agents": ["Data Loader", "Data Analyzer", "Report Writer"],
    "auto": false,
    "rounds": 1
  }
}

═══════════════════════════════════════════════
📘 FULL EXAMPLE — LOOPING SCRIPT BUILDER
(Shows: auto mode, script generation, subprocess, is_done, next_agent routing)
═══════════════════════════════════════════════
{
  "system_name": "Python Script Builder",
  "agents": [
    {
      "name": "Requirements Gatherer",
      "description": "Asks user what Python script they want built. Saves requirements as script_requirements.",
      "sub_agents": [
        {
          "prompt": "Ask user to describe the Python script they want built. Save the answer as script_requirements.",
          "code": "script_requirements = ask_user('Describe the Python script you want me to build:')\nif not script_requirements:\n    print('No requirements provided.')\n    is_done = True\nelse:\n    print(f'Requirements received: {script_requirements}')"
        }
      ]
    },
    {
      "name": "Script Generator",
      "description": "Reads script_requirements, generates a Python script using AI, strips markdown, saves to Desktop, runs it, saves output as script_output.",
      "sub_agents": [
        {
          "prompt": "Read script_requirements from memory. Use AI to generate a complete Python script. Strip any markdown fences. Save to Desktop as generated_script.py. Run it with subprocess and save output as script_output.",
          "code": "import re, os, sys, subprocess\\nraw = session.post('https://api.zygoflow.com/ai/analyze-text', json={'website_id': website_id, 'text': script_requirements, 'instruction': 'Write a complete self-contained Python script that does exactly what is described. Return ONLY raw Python code. No markdown, no backticks, no explanation.'}).json()\\ngenerated_script = raw.get('result', '')\\ngenerated_script = re.sub(r'^```(?:python)?\\n?', '', generated_script.strip())\\ngenerated_script = re.sub(r'\\n?```$', '', generated_script.strip())\\nif not generated_script:\\n    print('Error: AI failed to generate script.')\\n    is_done = True\\nelse:\\n    script_path = os.path.join(os.path.expanduser('~'), 'Desktop', 'generated_script.py')\\n    with open(script_path, 'w') as f:\\n        f.write(generated_script)\\n    print(f'Script saved to: {script_path}')\\n    result = subprocess.run([sys.executable, script_path], capture_output=True, text=True, timeout=180)\\n    script_output = result.stdout\\n    if result.stderr:\\n        script_output += '\\nERRORS:\\n' + result.stderr\\n    print(script_output)"
        }
      ]
    },
    {
      "name": "Quality Checker",
      "description": "Reads script_output from memory. If errors found, routes back to Script Generator. If clean, sets is_done = True.",
      "sub_agents": [
        {
          "prompt": "Read script_output from memory. Check if it contains errors. If errors exist, print them and set next_agent to 'Script Generator'. If clean, print success and set is_done = True.",
          "code": "if 'ERROR' in script_output.upper() or 'TRACEBACK' in script_output.upper():\\n    print(f'Errors detected:\\n{script_output}')\\n    print('Routing back to Script Generator to fix...')\\n    next_agent = 'Script Generator'\\nelse:\\n    print('Script ran successfully! Pipeline complete.')\\n    is_done = True"
        }
      ]
    }
  ],
  "pipeline": {
    "name": "Python Script Builder",
    "agents": ["Requirements Gatherer", "Script Generator", "Quality Checker"],
    "auto": true,
    "rounds": 1
  }
}

═══════════════════════════════════════════════
💬 CONVERSATION STYLE
═══════════════════════════════════════════════
- Be concise and friendly
- In DESIGN stage: propose clearly with emojis, name each agent, state its inputs and outputs
- In CONFIRM stage: summarize as a neat numbered list, ask "Shall I build this now?"
- In BUILD stage: return the full system JSON immediately — no extra questions, no placeholders
- ALWAYS return valid JSON — never raw text outside the JSON wrapper

═══════════════════════════════════════════════
⚡ RESPONSE FORMAT (ALWAYS return this exact structure)
═══════════════════════════════════════════════
Always return a JSON object with exactly these keys:
{
  "reply": "Your conversational message to the user",
  "stage": "designing" | "confirming" | "building" | "done",
  "system": null
}

When stage is "building":
- "system" must contain the full system JSON described in OUTPUT FORMAT above
- "reply" must be: "✅ Building your system now — saving all agents and pipeline..."
- Return immediately — do not ask any more questions

When stage is "done":
- "system" is null
- "reply" is a short confirmation message

Return ONLY this JSON object.
No markdown. No backticks. No extra text outside the JSON. Ever.
"""
        messages = []
        for msg in request.messages:
            role = msg.get("role", "user")
            if role == "system":
                continue  # skip — we pass system separately
            messages.append({
                "role": role,
                "content": msg.get("content", "")
            })

        resp = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=8000,
            system=system_prompt,  # same system_prompt string as GPT version
            messages=messages,
            temperature=0.4,
        )

        raw_result = resp.content[0].text.strip()

        # Claude might wrap in markdown fences even when told not to — strip them
        import re, json
        raw_result = re.sub(r'^```(?:json)?\n?', '', raw_result)
        raw_result = re.sub(r'\n?```$', '', raw_result)

        parsed = json.loads(raw_result.strip())

        # Track usage
        usage = getattr(resp, "usage", None)
        if usage and request.website_id and len(request.website_id) >= 32:
            await track_ai_usage(
                db=db,
                website_id=request.website_id,
                user_id=current_user.id,
                model="claude-sonnet-4-6",
                feature="architect",
                prompt_tokens=int(getattr(usage, "input_tokens", 0) or 0),
                completion_tokens=int(getattr(usage, "output_tokens", 0) or 0),
                meta={"stage": parsed.get("stage", "unknown")}
            )

        return {
            "reply": parsed.get("reply", ""),
            "stage": parsed.get("stage", "designing"),
            "system": parsed.get("system", None)
        }

    except Exception as e:
        print(f"Architect Claude Error: {e}")
        raise HTTPException(status_code=500, detail=str(e))


class GenerateComplexRequest(BaseModel):
    website_id: UUID
    prompt: str
    unique_class_name: str
 
 
@router.post("/generate-complex-element")
async def generate_complex_element(
    body: GenerateComplexRequest,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_active_user),
):
    """Two-pass generic pipeline: architect spec → provision persistence → build element."""
    try:
        website = await get_website(body.website_id, db)
        client, model, provider = get_ai_client(website)
 
        # ---- PASS 1: ARCHITECT ----
        lib_menu = ", ".join(f"{k} ({v['what']})" for k, v in COMPLEX_LIB_REGISTRY.items())
        spec = call_ai_json(
            client, model, provider,
            COMPLEX_ARCHITECT_PROMPT.replace("{LIB_MENU}", lib_menu),
            f'REQUEST: "{body.prompt}"',
            temperature=0.3, max_tokens=3000,
        )
        if not isinstance(spec, dict) or not spec.get("views"):
            raise HTTPException(500, "Architect produced no usable spec.")
 
        # ---- PROVISION PERSISTENCE ----
        schema_id = None
        schema_fields = []
        persistence = spec.get("persistence", "none")
 
        if persistence == "document":
            tname = spec.get("document_table_name") or f"{spec.get('element_name','App')} Documents"
            existing = (await db.execute(select(CustomDataSchema).where(
                CustomDataSchema.website_id == body.website_id,
                CustomDataSchema.name == tname))).scalars().first()
            doc_fields = [
                {"id": "title", "label": "Title", "type": "text"},
                {"id": "doc", "label": "Document", "type": "text"},
            ]
            if existing:
                schema_id, schema_fields = str(existing.schema_id), existing.fields
            else:
                ns = CustomDataSchema(website_id=body.website_id, name=tname, fields=doc_fields)
                db.add(ns)
                await db.commit()
                await db.refresh(ns)
                schema_id, schema_fields = str(ns.schema_id), doc_fields
 
        elif persistence == "rows":
            name_map = {}
            existing_all = (await db.execute(select(CustomDataSchema).where(
                CustomDataSchema.website_id == body.website_id))).scalars().all()
            for s in existing_all:
                name_map[s.name] = str(s.schema_id)
            # two-pass: non-relation fields first, then resolve related_table names
            raw_tables = spec.get("row_tables", []) or []
            for t in raw_tables:
                if not isinstance(t, dict) or not t.get("name") or t["name"] in name_map:
                    continue
                base_fields = [f for f in (t.get("fields") or []) if f.get("type") != "relation"]
                ns = CustomDataSchema(website_id=body.website_id, name=t["name"], fields=base_fields)
                db.add(ns)
                await db.flush()
                name_map[t["name"]] = str(ns.schema_id)
            await db.commit()
            for t in raw_tables:
                sid = name_map.get(t.get("name", ""))
                if not sid:
                    continue
                resolved = []
                for f in (t.get("fields") or []):
                    fd = dict(f)
                    if fd.get("type") == "relation":
                        target = name_map.get(fd.pop("related_table", ""))
                        if not target:
                            continue
                        fd["related_schema_id"] = target
                    resolved.append(fd)
                sch = await db.get(CustomDataSchema, UUID(sid))
                if sch:
                    sch.fields = resolved
                    flag_modified(sch, "fields")
            await db.commit()
            first = (spec.get("row_tables") or [{}])[0].get("name")
            if first and first in name_map:
                schema_id = name_map[first]
                sch = await db.get(CustomDataSchema, UUID(schema_id))
                schema_fields = sch.fields if sch else []
 
        # ---- PASS 2: BUILD ----
        libs = [
            {"name": n, "url": COMPLEX_LIB_REGISTRY[n]["url"], "global": COMPLEX_LIB_REGISTRY[n]["global"]}
            for n in (spec.get("libraries") or []) if n in COMPLEX_LIB_REGISTRY
        ]
        all_schemas = (await db.execute(select(CustomDataSchema).where(
            CustomDataSchema.website_id == body.website_id))).scalars().all()
        all_schemas_summary = [{"name": s.name, "schema_id": str(s.schema_id), "fields": s.fields}
                               for s in all_schemas]
 
        build_content = (
            f"SPEC:\n{json.dumps(spec, indent=2)}\n\n"
            f"LIBRARIES: {json.dumps(libs)}\n\n"
            f"UNIQUE_CLASS_NAME: `.{body.unique_class_name}`\n\n"
            f"SCHEMA_ID: {schema_id or 'none'}\n"
            f"SCHEMA_FIELDS: {json.dumps(schema_fields)}\n"
            f"EXISTING_SCHEMAS_ON_WEBSITE: {json.dumps(all_schemas_summary)}"
        )
        payload = generate_with_repair(
            client, model, provider, COMPLEX_BUILDER_PROMPT, build_content,
            max_tokens=20000,
        )
 
        # ---- ASSEMBLE (same shape as generate-ai-element) ----
        props = payload.get("properties", {}) or {}
        if schema_id:
            props["schema_id"] = schema_id
            props["schema_fields"] = schema_fields
        props["all_schemas"] = all_schemas_summary
        props["website_id"] = str(body.website_id)
        props["subdomain"] = website.subdomain
 
        return {
            "aiTemplate": f'<div class="{body.unique_class_name}">{payload.get("aiTemplate", "")}</div>'
                          if not payload.get("aiTemplate", "").strip().startswith(f'<div class="{body.unique_class_name}"')
                          else payload.get("aiTemplate", ""),
            "properties": props,
            "editableProps": payload.get("editableProps", []),
            "script": inject_runtime_lib(sanitize_injected_params(payload.get("script", ""))),
        }
 
    except HTTPException:
        raise
    except Exception as e:
        import traceback; traceback.print_exc()
        raise HTTPException(500, f"Complex element generation failed: {e}")

 
 
